"""小Q文档解析:PDF / 图片 / Word / PPT / Excel → 文本,或按分类体系抽字段汇总成 Excel。

对齐官方 skill(quickbi-smartq-chat · 文档解析)的两种模式:

    text      纯文本提取:逐个文件取出正文,原样给出
    extract   分类 + 字段提取:按 resources/document_classification.zh.md(10 大类 37 种)
              归类、抽字段,一类一张 Sheet,首页「汇总」,产出一份 xlsx

取文本的顺序也与官方一致 —— **本地优先、远程兜底**:本地能读出正文的(docx / 带文字层的 PDF /
表格 / txt)不必上传;扫描件、图片、老版 doc、ppt 这类本地读不出字的,走 Quick BI 的
远程 OCR(document/upload → document/task 轮询)。

这一层只依赖 QuickBIClient 与两个回调(llm / save_output),不碰数据库:
抽字段要调大模型、产出的 Excel 要落成上传文件,这两件事由 service 层注入。
"""
from __future__ import annotations

import io
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, AsyncIterator, Awaitable, Callable, Sequence

from cogni_agent.tool_center_service.quickbi import filesupport as fs
from cogni_agent.tool_center_service.quickbi.client import QuickBIClient
from cogni_agent.tool_center_service.quickbi.errors import QuickBIError
from cogni_agent.tool_center_service.quickbi.events import make_chart

URI_UPLOAD = "/openapi/v2/document/upload"
URI_TASK = "/openapi/v2/document/task"

# 远程 OCR 轮询:首次 3 秒,每次 ×1.5,封顶 10 秒,总共最多等 120 秒(与官方脚本一致)
POLL_INTERVAL = 3.0
POLL_BACKOFF = 1.5
POLL_MAX_INTERVAL = 10.0
POLL_MAX_WAIT = 120.0
# 间隔被设成 0(测试 / 加速)时靠次数兜底,免得上游一直 processing 就死循环
POLL_MAX_TIMES = 60

# 本地读出来的字少于这个数,视为「没读出来」(扫描件 PDF 的文字层常常只有页码)
MIN_LOCAL_CHARS = 20
# 文本模式每个文件最多给多少字;抽字段时每个文件喂给模型多少字、一共多少字
MAX_TEXT_PER_FILE = 20_000
MAX_EXTRACT_PER_FILE = 8_000
MAX_EXTRACT_TOTAL = 60_000

UNRECOGNIZED = "未识别"
PARSE_FAILED = "解析失败"

# 提这些词就是要「整理成表」,否则只提正文
_EXTRACT_KW = ("提取", "抽取", "字段", "台账", "汇总", "整理成表", "整理成excel", "导出excel", "生成excel",
               "excel", "表格", "分类", "归档", "识别发票", "录入", "extract", "classify")

# 汇总页的分类组(与官方 generate_excel.py 的 GROUP_MAP 同口径:子类型 → 分类组)
GROUP_MAP = {
    **dict.fromkeys(("增值税发票", "银行回单", "银行对账单", "费用报销单", "合同/协议", "合同协议", "收据",
                     "财务报表/审计报告", "财务报表", "其他金额类文档"), "A.财务与税务"),
    **dict.fromkeys(("劳动合同", "简历", "绩效评估", "薪资单", "离职证明", "培训记录", "员工档案"), "B.人力资源"),
    **dict.fromkeys(("采购合同", "入库单/出库单", "入库单", "质检报告", "供应商资质", "物流订单",
                     "报价单/询价单"), "C.供应链与采购"),
    **dict.fromkeys(("营业执照", "身份证", "护照", "保密协议", "资质证书", "公文/通知", "公文通知"), "D.行政与法务"),
    **dict.fromkeys(("病历", "处方单", "检验/检查报告", "检验检查报告", "体检报告"), "E.医疗"),
    **dict.fromkeys(("保单", "理赔申请", "理赔结案通知"), "F.保险"),
    **dict.fromkeys(("运单", "提单", "报关单"), "G.物流"),
    **dict.fromkeys(("系统日志", "漏洞扫描报告", "服务器监控报表"), "H.技术与运维"),
    **dict.fromkeys(("客诉/工单记录", "客诉工单记录", "销售报价单", "售后退换货单"), "I.客服与销售"),
    **dict.fromkeys(("招投标文件", "政务审批单", "合规审计报告"), "J.政务与合规"),
    UNRECOGNIZED: UNRECOGNIZED,
    PARSE_FAILED: PARSE_FAILED,
}

Llm = Callable[[list[dict]], Awaitable[str]]
SaveOutput = Callable[[str, bytes], Awaitable[dict]]
Sleep = Callable[[float], Awaitable[Any]]


def status(stage: str, message: str, **extra) -> dict:
    return {"type": "status", "stage": stage, "message": message, **extra}


def pick_mode(question: str, mode: str | None = None) -> str:
    """text / extract。显式指定优先;auto(或不填)按问句判断。"""
    if mode in ("text", "extract"):
        return mode
    q = (question or "").lower()
    return "extract" if any(k in q for k in _EXTRACT_KW) else "text"


# ═══════════════════════════ 取文本:本地 ═══════════════════════════

def local_text(name: str, content: bytes) -> str | None:
    """本地能读出正文就返回,读不出(格式不支持 / 缺库 / 扫描件)返回 None,交给远程 OCR。"""
    from cogni_agent.tool_center_service import filereader as fr

    s = fs.suffix(name)
    try:
        if s in fs.INSIGHT_SUFFIXES:
            text = fs.to_markdown(name, content)
        elif s in (".docx", ".pdf", ".txt", ".md"):
            text = fr.read(name, content).get("text") or ""
        else:
            # 图片 / doc / ppt / pptx:本地 OCR 依赖系统 tesseract,线上多半没有,
            # 而且识别率远不如 Quick BI 的 OCR —— 直接走远程
            return None
    except Exception:  # noqa: BLE001 - 本地读不了就交给远程,不在这里报错
        return None
    return text if len(re.sub(r"\s+", "", text or "")) >= MIN_LOCAL_CHARS else None


# ═══════════════════════════ 取文本:远程 OCR ═══════════════════════════

async def upload_document(qbi: QuickBIClient, name: str, content: bytes) -> str:
    """document/upload:multipart 只带 file,签名参数为空。返回 taskId。"""
    files = {"file": (name, content, fs.MIME.get(fs.suffix(name), "application/octet-stream"))}
    data = await qbi.request("POST", URI_UPLOAD, files=files, sign_params=None, timeout=120)
    if isinstance(data, str):
        try:
            data = json.loads(data)
        except json.JSONDecodeError:
            data = {}
    task_id = (data or {}).get("taskId") if isinstance(data, dict) else None
    if not task_id:
        raise QuickBIError(f"文档上传成功但未返回 taskId:{str(data)[:300]}", kind="protocol")
    return str(task_id)


async def remote_text(qbi: QuickBIClient, name: str, content: bytes, *, sleep: Sleep,
                      interval: float = POLL_INTERVAL, max_wait: float = POLL_MAX_WAIT) -> AsyncIterator[dict]:
    """上传 + 轮询。产出 status 事件,最后产出 {"type": "doc_text", "text": ...}。失败抛 QuickBIError。"""
    task_id = await upload_document(qbi, name, content)
    yield status("ocr", f"「{name}」已提交 Quick BI 识别,等待结果…", task_id=task_id)
    waited, wait, times = 0.0, interval, 0
    while True:
        data = await qbi.request("GET", URI_TASK, params={"taskId": task_id}, timeout=60)
        data = data if isinstance(data, dict) else {}
        st = str(data.get("status") or "").lower()
        if data.get("completed") or st in ("success", "failed", "not_supported"):
            if st == "success":
                yield {"type": "doc_text", "text": str(data.get("parsedText") or "")}
                return
            if st == "not_supported":
                raise QuickBIError(f"Quick BI 不支持识别「{name}」这种文件", kind="input", code="DOC_NOT_SUPPORTED")
            raise QuickBIError(f"「{name}」识别失败:{data.get('errorMessage') or st or '未知原因'}",
                               kind="business", code="DOC_PARSE_FAILED")
        times += 1
        if waited >= max_wait or times >= POLL_MAX_TIMES:
            raise QuickBIError(f"「{name}」识别超时(已等待 {int(waited)} 秒)", kind="network", code="DOC_TIMEOUT")
        await sleep(wait)
        waited += wait
        wait = min(wait * POLL_BACKOFF, POLL_MAX_INTERVAL) if wait else 0


# ═══════════════════════════ 抽字段 ═══════════════════════════

_CLASSIFICATION: str | None = None


def classification_reference() -> str:
    global _CLASSIFICATION
    if _CLASSIFICATION is None:
        _CLASSIFICATION = (Path(__file__).parent / "resources" / "document_classification.zh.md").read_text("utf-8")
    return _CLASSIFICATION


def extract_messages(question: str, docs: Sequence[tuple[str, str]]) -> list[dict]:
    """分类 + 抽字段的提示词。docs = [(文件名, 正文)]。"""
    system = (
        "你是文档结构化提取员。按下面的《文档分类体系》给每份文档归类,并按该子类的字段定义提取字段值。\n"
        "规则:\n"
        "1. 只填文档里**真实出现**的值,找不到的字段留空字符串,禁止推测、补全或编造;\n"
        "2. 每个子类一张表,表名用分类体系里的中文 Sheet 名;表头用中文,第一列固定为「源文件名」;\n"
        "3. 无法归入 37 种子类型、或 OCR 质量差无法确定的,归入「未识别」表,"
        "表头固定为 [源文件名, 疑似类型, 可提取的关键字段, 说明];\n"
        "4. 一份文档里有多条明细(如对账单多笔流水)时,按分类体系的约定处理,不要拆成多份文档;\n"
        "5. 只输出一个 JSON 对象,不要输出任何解释或 Markdown 代码块,格式:\n"
        '{"extraction_data": {"<表名>": {"headers_cn": ["源文件名", ...], "rows": [["文件名", ...], ...]}}}\n\n'
        "《文档分类体系》\n" + classification_reference()
    )
    budget = MAX_EXTRACT_TOTAL
    parts = []
    for name, text in docs:
        chunk = (text or "")[:min(MAX_EXTRACT_PER_FILE, max(budget, 0))]
        budget -= len(chunk)
        parts.append(f"=== 文件:{name} ===\n{chunk or '(空)'}")
    user = (f"用户要求:{question}\n\n" if question else "") + "待处理文档:\n\n" + "\n\n".join(parts)
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


def parse_extraction(text: str) -> dict[str, dict]:
    """模型输出 → {表名: {headers_cn, rows}}。容忍外面包了 ```json 代码块或前后有话。"""
    s = (text or "").strip()
    fence = re.search(r"```(?:json)?\s*(.*?)```", s, re.S)
    if fence:
        s = fence.group(1).strip()
    if not s.startswith("{"):
        start, end = s.find("{"), s.rfind("}")
        s = s[start:end + 1] if 0 <= start < end else ""
    try:
        data = json.loads(s)
    except (json.JSONDecodeError, ValueError) as exc:
        raise QuickBIError(f"字段提取结果不是合法 JSON:{(text or '')[:200]}", kind="protocol",
                           code="DOC_EXTRACT_INVALID") from exc
    sheets = data.get("extraction_data", data) if isinstance(data, dict) else None
    if not isinstance(sheets, dict):
        raise QuickBIError("字段提取结果缺少 extraction_data", kind="protocol", code="DOC_EXTRACT_INVALID")
    out: dict[str, dict] = {}
    for name, sheet in sheets.items():
        if not isinstance(sheet, dict):
            continue
        headers = [str(h) for h in sheet.get("headers_cn") or []]
        rows = [[("" if v is None else v) for v in r] for r in sheet.get("rows") or [] if isinstance(r, list)]
        if headers:
            out[_sheet_name(str(name))] = {"headers_cn": headers, "rows": rows}
    return out


def _sheet_name(name: str) -> str:
    # Excel 的 Sheet 名:≤31 字符,不能含 []:*?/\
    return re.sub(r"[\[\]:*?/\\]", "_", name).strip()[:31] or UNRECOGNIZED


# ═══════════════════════════ 产出 Excel ═══════════════════════════

def build_excel(sheets: dict[str, dict], *, scan_time: str, total_files: int) -> bytes:
    """一类一张 Sheet,首页「汇总」。格式与官方 generate_excel.py 一致:蓝底白字粗体表头、
    细边框、冻结首行、自动筛选、按中文宽度自动列宽。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fill = PatternFill("solid", fgColor="4472C4")
    font = Font(bold=True, color="FFFFFF", size=11)
    side = Side("thin")
    border = Border(left=side, right=side, top=side, bottom=side)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")

    def width(v) -> int:
        return sum(2 if ord(c) > 0x7F else 1 for c in str(v or ""))

    def write(ws, headers: list, rows: list) -> None:
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill, cell.font, cell.alignment, cell.border = fill, font, center, border
        for r, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border, cell.alignment = border, wrap
        for c in range(1, len(headers) + 1):
            w = max([width(headers[c - 1])] + [width(row[c - 1]) for row in rows if c - 1 < len(row)])
            ws.column_dimensions[get_column_letter(c)].width = min(w + 2, 50)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    summary = wb.active
    summary.title = "汇总"
    write(summary, ["分类组", "子类型", "文件数量", "提取字段"],
          [[GROUP_MAP.get(name, "其他"), name, len({str(r[0]) for r in s["rows"] if r}),
            ", ".join(s["headers_cn"][1:])] for name, s in sheets.items()])
    meta_row = summary.max_row + 2
    for i, (k, v) in enumerate((("扫描时间", scan_time), ("文件总数", total_files))):
        summary.cell(row=meta_row + i, column=1, value=k).font = Font(bold=True)
        summary.cell(row=meta_row + i, column=2, value=v)
    for name, s in sheets.items():
        write(wb.create_sheet(title=name), s["headers_cn"], s["rows"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheet_charts(sheets: dict[str, dict], *, max_rows: int = 200) -> list[dict]:
    """每张 Sheet 一张明细表(table),页面上直接看、给模型的文本里也带得出来。"""
    out = []
    for name, s in sheets.items():
        fields = [{"name": h, "type": "string", "role": "dimension"} for h in s["headers_cn"]]
        out.append(make_chart(title=name, fields=fields, rows=s["rows"][:max_rows], chart_type="table",
                              source="doc_parse"))
    return out


# ═══════════════════════════ 编排 ═══════════════════════════

async def run_doc_parse(qbi: QuickBIClient, *, question: str, files: Sequence[tuple[str, bytes]],
                        mode: str | None = None, llm: Llm | None = None, save_output: SaveOutput | None = None,
                        sleep: Sleep, interval: float = POLL_INTERVAL,
                        now: datetime | None = None) -> AsyncIterator[dict]:
    """files = [(文件名, 内容)]。逐个取文本(本地优先、远程兜底),再按模式给正文或抽字段出 Excel。"""
    if not files:
        raise QuickBIError("文档解析需要先上传文件(file_id / file_ids)", kind="input")
    mode = pick_mode(question, mode)
    yield {"type": "meta", "doc_mode": mode, "doc_files": [n for n, _ in files]}
    docs: list[tuple[str, str]] = []
    failed: list[tuple[str, str]] = []
    for name, content in files:
        s = fs.suffix(name)
        if s not in fs.DOC_PARSE_SUFFIXES:
            failed.append((name, f"不支持的文件类型 {s or '(无后缀)'}"))
            yield status("parse", f"「{name}」不是文档解析支持的类型,已跳过")
            continue
        text = local_text(name, content)
        if text is not None:
            yield status("parse", f"「{name}」本地解析完成({len(text)} 字)", via="local")
            docs.append((name, text))
            continue
        yield status("upload", f"正在上传「{name}」到 Quick BI 识别…", via="remote")
        try:
            async for e in remote_text(qbi, name, content, sleep=sleep, interval=interval):
                if e["type"] == "doc_text":
                    docs.append((name, e["text"]))
                    yield status("parse", f"「{name}」识别完成({len(e['text'])} 字)", via="remote")
                else:
                    yield e
        except QuickBIError as exc:
            if exc.kind == "config":
                raise                     # 凭证/权限问题,换下一个文件也一样,直接报
            failed.append((name, str(exc)))
            yield status("parse_failed", str(exc), file_name=name)
    if not docs:
        raise QuickBIError("没有可用的文档内容:" + ";".join(f"{n}:{m}" for n, m in failed),
                           kind="business", code="DOC_PARSE_FAILED")

    if mode == "text":
        parts = []
        for name, text in docs:
            body = text.strip() or "(没有识别到文字)"
            if len(body) > MAX_TEXT_PER_FILE:
                body = body[:MAX_TEXT_PER_FILE] + f"\n\n(正文共 {len(text)} 字,仅展示前 {MAX_TEXT_PER_FILE} 字)"
            parts.append(f"## {name}\n\n{body}")
        if failed:
            parts.append("## 未能解析\n\n" + "\n".join(f"- {n}:{m}" for n, m in failed))
        yield {"type": "text", "section": "answer", "delta": "\n\n".join(parts)}
        return

    if llm is None:
        raise QuickBIError("字段提取需要大模型,当前环境未配置", kind="config")
    yield status("extract", f"正在对 {len(docs)} 份文档分类并提取字段…")
    sheets = parse_extraction(await llm(extract_messages(question, docs)))
    if failed:
        sheets[PARSE_FAILED] = {"headers_cn": ["源文件名", "错误信息", "文件类型"],
                                "rows": [[n, m, fs.suffix(n)] for n, m in failed]}
    for e in [{"type": "chart", "chart": c} for c in sheet_charts(sheets)]:
        yield e
    n = now or datetime.now()
    summary = [f"共 {len(files)} 份文件,成功解析 {len(docs)} 份。"]
    summary += [f"- {name}:{len(s['rows'])} 行" for name, s in sheets.items()]
    unknown = sheets.get(UNRECOGNIZED, {}).get("rows") or []
    if unknown:
        summary.append(f"有 {len(unknown)} 份文档未能归入预定义类型。如需按自定义类型提取,"
                       "请确认要提取的字段(至少 5 个),再按「Custom_类型名」重新提取。")
    if save_output is not None:
        name = f"doc_scan_result_{n.strftime('%Y%m%d_%H%M%S')}.xlsx"
        content = build_excel(sheets, scan_time=n.strftime("%Y-%m-%d %H:%M:%S"), total_files=len(files))
        saved = await save_output(name, content)
        yield {"type": "meta", "output_file": {"id": saved.get("id"), "file_name": name, "size": len(content)}}
        summary.append(f"已生成汇总 Excel「{name}」(file_id={saved.get('id')}),"
                       "可以直接拿它继续做文件问数或解读。")
    yield {"type": "text", "section": "answer", "delta": "\n".join(summary)}
