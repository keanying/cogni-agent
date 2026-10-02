# -*- coding: utf-8 -*-
"""读用户上传的文件。

**这一组守的核心判断是:表格类文件不该变成文本。**

一份三千行的 Excel 转成 Markdown 表格是四十多万字,塞不进上下文;截断之后模型看到的是
半张表,然后一本正经地基于半张表下结论 —— 不报错,也没人会去核对它是不是只读了一半。
所以表格读进来变成**数据集句柄**,和接口取回来的大数据走完全同一条路。

其余几条守的都是「读错了但不报错」:
- 中文 CSV 现场多半是 GBK,一律按 UTF-8 读会整篇变乱码,而乱码不会抛异常;
- Word 的关键信息大半在表格里,只抽段落等于把最该读的那部分丢了;
- 扫描件 PDF 抽不出文字,不说破的话会被当成「这文件是空的」;
- 多 sheet 的 Excel 只读第一个而不说,模型会把它当成整个文件;
- OCR 没装中文包时硬认中文,会认出一堆无意义字母 —— 而模型会认真分析那堆字母。
"""
from __future__ import annotations

import io

import pytest

from cogni_agent.tool_center_service import filereader as fr


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════
# 一、表格 → 数据集,不是文本
# ══════════════════════════════════════════════════════════════════════

def test_excel_becomes_rows_not_text():
    got = fr.read("客源.xlsx", _xlsx({"客源": [["城市", "热度"], ["福建省", 62.12], ["江西省", 8.2]]}))
    assert got["kind"] == "table"
    assert got["rows"] == [{"城市": "福建省", "热度": 62.12}, {"城市": "江西省", "热度": 8.2}]


def test_all_sheet_names_are_reported():
    """只读第一个 sheet 而不说,模型会把它当成整个文件 —— 而另外几个表里
    可能正好是用户问的那部分。"""
    got = fr.read("多表.xlsx", _xlsx({"客源": [["城市"], ["福建省"]],
                                      "票务": [["票种"], ["全价票"]],
                                      "说明": [["备注"], ["内部"]]}))
    assert got["sheets"] == ["客源", "票务", "说明"]


def test_a_specific_sheet_can_be_read():
    got = fr.read("多表.xlsx", _xlsx({"客源": [["城市"], ["福建省"]],
                                      "票务": [["票种"], ["全价票"]]}), sheet="票务")
    assert got["rows"] == [{"票种": "全价票"}]


def test_blank_rows_inside_excel_are_skipped():
    """Excel 里中间空一行太常见了。不跳过的话会多出一行全 None,
    而它会被算进合计的分母 —— 均值就这么悄悄错了。"""
    got = fr.read("a.xlsx", _xlsx({"s": [["城市", "热度"], ["福建省", 62.12],
                                         [None, None], ["江西省", 8.2]]}))
    assert len(got["rows"]) == 2


def test_header_only_excel_says_so():
    with pytest.raises(fr.FileReadError) as e:
        fr.read("a.xlsx", _xlsx({"s": [["城市", "热度"]]}))
    assert "没有数据行" in str(e.value)


def test_chinese_csv_in_gbk_is_decoded_correctly():
    """**中文 CSV 现场多半是 GBK。**

    一律按 UTF-8 读会整篇变乱码,而解码失败不会抛异常(errors='replace' 会给一串问号),
    模型会对着那堆问号认真分析,并给出一份看起来很完整的结论。
    """
    got = fr.read("a.csv", "城市,热度\n福建省,62.12\n".encode("gbk"))
    assert got["rows"] == [{"城市": "福建省", "热度": "62.12"}]


def test_tsv_and_semicolon_csv_are_sniffed():
    assert fr.read("a.tsv", "城市\t热度\n福建省\t62\n".encode())["rows"][0]["热度"] == "62"
    assert fr.read("a.csv", "城市;热度\n福建省;62\n".encode())["rows"][0]["热度"] == "62"


def test_row_cap_is_enforced():
    """几万行还能读,但要有上限 —— 内存是所有任务共用的。"""
    big = "a,b\n" + "\n".join(f"{i},{i}" for i in range(fr.MAX_TABLE_ROWS + 500))
    assert len(fr.read("x.csv", big.encode())["rows"]) == fr.MAX_TABLE_ROWS


# ══════════════════════════════════════════════════════════════════════
# 二、文档
# ══════════════════════════════════════════════════════════════════════

def test_word_tables_are_kept_as_markdown():
    """**Word 的关键信息大半在表格里。**

    合同的金额、报表的指标,都在表格中;只抽段落的话正文读到了、数字全丢了,
    而模型不会知道自己少读了什么。
    """
    import docx

    d = docx.Document()
    d.add_paragraph("九龙潭 2026 上半年客源分析")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "城市", "占比"
    t.cell(1, 0).text, t.cell(1, 1).text = "福建省", "62.12%"
    b = io.BytesIO()
    d.save(b)

    text = fr.read("x.docx", b.getvalue())["text"]
    assert "九龙潭 2026 上半年客源分析" in text
    assert "| 福建省 | 62.12% |" in text, "表格必须留下来,而且是能读的结构"


def test_empty_word_points_at_ocr():
    """整篇是图片的 Word(扫描件)抽不出字。说清下一步,别让它变成「这文件是空的」。"""
    import docx

    b = io.BytesIO()
    docx.Document().save(b)
    with pytest.raises(fr.FileReadError) as e:
        fr.read("x.docx", b.getvalue())
    assert "OCR" in str(e.value)


def test_plain_text_and_json_are_read_as_documents():
    assert "hello" in fr.read("a.txt", b"hello world")["text"]
    assert fr.read("a.json", b'{"a":1}')["kind"] == "doc"


# ══════════════════════════════════════════════════════════════════════
# 三、拒绝与降级:错要说清是哪一种错
# ══════════════════════════════════════════════════════════════════════

def test_unsupported_type_lists_what_is_supported():
    with pytest.raises(fr.FileReadError) as e:
        fr.read("a.exe", b"MZ")
    msg = str(e.value)
    assert ".xlsx" in msg and ".pdf" in msg, "要告诉用户支持什么,而不是只说不支持"


def test_oversized_file_is_refused_with_a_reason():
    with pytest.raises(fr.FileReadError) as e:
        fr.read("a.csv", b"x" * (fr.MAX_FILE_BYTES + 1))
    assert "上限" in str(e.value)


def test_empty_file_is_refused():
    with pytest.raises(fr.FileReadError):
        fr.read("a.csv", b"")


def test_kind_of_classifies_every_supported_suffix():
    for s in fr.TABLE_SUFFIXES:
        assert fr.kind_of(f"x{s}") == "table"
    for s in fr.DOC_SUFFIXES:
        assert fr.kind_of(f"x{s}") == "doc"
    for s in fr.IMAGE_SUFFIXES:
        assert fr.kind_of(f"x{s}") == "image"
    assert fr.kind_of("x.zip") == "unsupported"


def test_ocr_refuses_to_fake_chinese_with_an_english_model():
    """**没装中文包时不能硬认中文。**

    退回英文包的话,一张中文截图会认出一堆毫无意义的字母 —— 而 OCR 不会报错,
    模型会认真分析那堆字母,给出一份完全凭空的结论。
    """
    from PIL import Image

    img = Image.new("RGB", (200, 50), "white")
    b = io.BytesIO()
    img.save(b, "PNG")
    try:
        import pytesseract
        langs = set(pytesseract.get_languages(config=""))
    except Exception:                          # pragma: no cover
        pytest.skip("环境没有 tesseract")
    if "chi_sim" in langs:
        pytest.skip("这个环境装了中文包,测不到降级路径")
    with pytest.raises(fr.FileReadError) as e:
        fr.read("s.png", b.getvalue(), lang="chi_sim")
    assert "语言包" in str(e.value) and "不要拿英文包硬认中文" in str(e.value)


def test_ocr_reads_english_text():
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (420, 70), "white")
    ImageDraw.Draw(img).text((10, 25), "Fujian 62.12 percent", fill="black")
    b = io.BytesIO()
    img.save(b, "PNG")
    try:
        text = fr.read("s.png", b.getvalue(), lang="eng")["text"]
    except fr.FileReadError as exc:            # pragma: no cover
        pytest.skip(f"环境 OCR 不可用:{exc}")
    assert "62.12" in text


# ══════════════════════════════════════════════════════════════════════
# 四、接上内置工具:表格 → 句柄 → 统计/画图
# ══════════════════════════════════════════════════════════════════════

def test_uploaded_table_flows_into_the_same_toolchain():
    """上传的表格和接口取回来的大数据必须走同一条路 —— 否则要维护两套统计能力,
    而两套迟早长歪,歪掉的那套通常没人天天用。"""
    from cogni_agent.runtime_service.engine import databuf
    from cogni_agent.tool_center_service import builtin_tools as bi

    task = "task-file-1"
    databuf.clear(task)
    rows = fr.read("客源.xlsx", _xlsx({"s": [["城市", "热度"],
                                            ["福建省", 62.12], ["江西省", 8.2],
                                            ["上海市", 6.79]]}))["rows"]
    h = databuf.stash(task, "上传文件:客源.xlsx", rows)["handle"]

    stats = bi.builtin_table_stats(dataset=h, columns="热度", _task_id=task)
    assert stats["columns"][0]["sum"] == 77.11
    chart = bi.builtin_chart(dataset=h, chart_type="bar", label_column="城市",
                             value_columns="热度", _task_id=task)
    assert chart.get("ok") is True and chart["points"] == 3
    databuf.clear(task)


def test_read_file_tool_is_registered_and_grouped():
    from cogni_agent.tool_center_service import builtin_tools as bi
    from cogni_agent.tool_center_service.routes import _BUILTIN_GROUPS

    assert "sys_read_file" in bi.BUILTIN
    assert _BUILTIN_GROUPS.get("sys_read_file")
    desc = bi.BUILTIN["sys_read_file"]["description"]
    assert "句柄" in desc, "工具说明必须讲清表格走句柄,否则模型会要求把明细贴出来"
