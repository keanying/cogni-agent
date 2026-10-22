"""Quick BI 集成模块 · 接口级集成测试(真实 PostgreSQL + 模拟 Quick BI)。

覆盖:连接(粘贴 skill 配置/加密/自检)→ 仪表板登记 → 文件上传 →
三个场景(固定问题 / 对话 / 报告)→ 开放平台 API Key 调用 → 一键生成 quickbi 工具 →
执行引擎解析(resolver)并经 invoke_tool 真正执行。

数据库:沿用 tests/conftest.py 的测试库({POSTGRES_DB}_test);只建本用例需要的表。
Quick BI:client.default_transport_factory 指向 mock_quickbi_server(不访问外网)。
"""
from __future__ import annotations

import asyncio
import functools
import io
import json

import httpx
import pytest
from cryptography.fernet import Fernet
from fastapi import FastAPI
from sqlalchemy import text

from cogni_agent.common import db as common_db
from cogni_agent.common import schema as _all_models  # noqa: F401  注册全部模型(建表用)
from cogni_agent.common.db import Base
from cogni_agent.common.errors import ZhilvError
from cogni_agent.common.middleware import ApiResponseMiddleware
from cogni_agent.config import get_settings
from cogni_agent.gateway_service.deps import CurrentUser, get_current_user
from cogni_agent.tool_center_service.quickbi import client as qclient
from cogni_agent.tool_center_service.quickbi import service as svc
from cogni_agent.tool_center_service.quickbi.models import QbiConnection
from cogni_agent.tool_center_service.quickbi.routes_admin import router as admin_router
from cogni_agent.tool_center_service.quickbi.routes_open import router as open_router
from tests.quickbi import mock_quickbi_server as mock

ADMIN = CurrentUser(user_id="u-admin", scenic_id=None, role="super_admin")
TOKEN3RD = ("https://bi.aliyuncs.com/token3rd/dashboard/view/pc.htm?pageId=88c9e5de-2d4d-4ba9-a045-6f9b71cf0c66"
            "&accessTicket=807cc156-9109-4dfc-b4d7-ce5c670ba566&dd_orientation=auto")


def run_async(fn):
    """与 pytest-asyncio / anyio 都不冲突的写法:同步测试函数内 asyncio.run。"""
    @functools.wraps(fn)
    def wrapper(*a, **kw):
        return asyncio.run(fn(*a, **kw))
    return wrapper


def _tables():
    return [t for t in Base.metadata.sorted_tables
            if not any(type(c.type).__module__.startswith("pgvector") for c in t.columns)]


@pytest.fixture(autouse=True)
def _env(monkeypatch):
    # 凭证加密密钥(测试专用)+ 试用凭证 = 模拟服务的 AK/SK
    from cogni_agent.tool_center_service.mcp import vault

    settings = get_settings()
    tc = dict(settings.raw.get("tool_center") or {})
    if not tc.get("credential_encrypt_key"):
        tc["credential_encrypt_key"] = Fernet.generate_key().decode()
        monkeypatch.setitem(settings.raw, "tool_center", tc)
        monkeypatch.setattr(vault, "_fernet_singleton", None)
    q = dict(settings.raw.get("quickbi") or {})
    q.update({"trial_api_key": mock.API_KEY, "trial_api_secret": mock.API_SECRET,
              "report_poll_seconds": 0, "snapshot_poll_seconds": 0,
              "default_connection": {"user_token": ""}})
    monkeypatch.setitem(settings.raw, "quickbi", q)
    monkeypatch.setattr(qclient, "default_transport_factory", lambda: httpx.ASGITransport(app=mock.app))
    mock.STATE.__init__()

    async def reset():
        from cogni_agent.user_center_service.models import User

        engine = common_db.get_engine()
        async with engine.begin() as conn:
            await conn.run_sync(lambda c: Base.metadata.drop_all(c, tables=_tables()))
            await conn.run_sync(lambda c: Base.metadata.create_all(c, tables=_tables()))
        async with common_db.get_sessionmaker()() as db:   # 工具所有者需要真实用户行(执行时按所有者校验)
            db.add(User(id=ADMIN.user_id, name="测试超管", account="qbi-admin", phone="13800000000",
                        password_hash="x", scenic_id="", role="super_admin"))
            await db.commit()
    asyncio.run(reset())
    yield


async def _error_handler(request, exc: ZhilvError):
    """与 main.py 的 zhilv_error_handler 同口径(不直接 import main,避免拉起整个应用依赖)。"""
    from fastapi.responses import JSONResponse

    from cogni_agent.common.schemas import ApiResponse
    return JSONResponse(status_code=exc.status_code,
                        content=ApiResponse.fail(code=exc.error_code or 1000, message=str(exc)).model_dump())


def _app(user: CurrentUser = ADMIN) -> FastAPI:
    app = FastAPI()
    app.add_middleware(ApiResponseMiddleware)
    app.add_exception_handler(ZhilvError, _error_handler)
    app.include_router(admin_router, prefix="/api")
    app.include_router(open_router)
    app.dependency_overrides[get_current_user] = lambda: user
    return app


def _client(app: FastAPI) -> httpx.AsyncClient:
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=30)


def data(resp: httpx.Response):
    body = resp.json()
    assert resp.status_code == 200 and body["code"] == 0, body
    return body["data"]


def sse_events(resp: httpx.Response) -> list[dict]:
    out = []
    for block in resp.text.replace("\r\n", "\n").split("\n\n"):
        lines = [ln[5:].strip() for ln in block.split("\n") if ln.startswith("data:")]
        if lines:
            out.append(json.loads("\n".join(lines)))
    return out


async def _setup_connection(c: httpx.AsyncClient) -> dict:
    # 用户给的配置:api_key/api_secret 为空,只有 user_token → 自动走试用凭证
    paste = "api_key:\napi_secret:\nuser_token:deadbeefcafe4f0e8a1b2c3d4e5f6071\nserver_domain:https://bi.aliyun.com"
    conn = data(await c.post("/api/admin/quickbi/connections/paste", json={"text": paste, "name": "羚羊QuickBI"}))
    return conn


# ─────────────────────────────────────────────────────────────────────

@run_async
async def test_connection_paste_encrypt_and_check():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        assert conn["use_trial"] and conn["trial_active"] and conn["is_default"]
        assert conn["missing"] == []
        assert conn["user_token"] == "dead****6071"         # 脱敏
        assert "deadbeefcafe4f0e8a1b2c3d4e5f6071" not in json.dumps(conn)
        # 数据库里是密文
        async with common_db.get_sessionmaker()() as db:
            row = await db.get(QbiConnection, conn["id"])
            assert row.user_token_cipher and mock.USER_TOKEN not in row.user_token_cipher
        chk = data(await c.post(f"/api/admin/quickbi/connections/{conn['id']}/check"))
        assert chk["ok"] and chk["cube_count"] == 2 and chk["trial"]
        cubes = data(await c.get(f"/api/admin/quickbi/connections/{conn['id']}/cubes", params={"keyword": "客流"}))
        assert cubes["items"] == [{"cube_id": mock.CUBE_TRAFFIC, "cube_name": "景区客流统计"}]
        lst = data(await c.get("/api/admin/quickbi/connections"))["items"]
        assert [x["status"] for x in lst] == ["ok"]


@run_async
async def test_bad_credentials_surface_clearly():
    async with _client(_app()) as c:
        conn = data(await c.post("/api/admin/quickbi/connections", json={
            "name": "错误密钥", "api_key": mock.API_KEY, "api_secret": "wrong", "user_token": mock.USER_TOKEN}))
        chk = data(await c.post(f"/api/admin/quickbi/connections/{conn['id']}/check"))
        assert chk["ok"] is False and "403" in chk["message"]
        r = await c.post("/api/admin/quickbi/run", json={"connection_id": conn["id"], "question": "销售额",
                                                          "target": {"cube_id": mock.CUBE_SALES}, "stream": False})
        run = data(r)
        assert run["status"] == "failed" and "signature" in run["error"]


@run_async
async def test_direct_run_stream_and_sync():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        r = await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dataset_qa", "question": "各平台销售额",
            "target": {"cube_id": mock.CUBE_SALES}, "stream": True})
        assert r.headers["content-type"].startswith("text/event-stream")
        evs = sse_events(r)
        assert evs[0]["type"] == "run" and evs[-1]["type"] == "done"
        assert evs[-1]["status"] == "success" and evs[-1]["chart_count"] == 1
        assert any(e["type"] == "chart" for e in evs)
        run = data(await c.get(f"/api/admin/quickbi/runs/{evs[0]['run_id']}"))
        assert run["result"]["charts"][0]["rows"][0]["销售额(SUM)"] == 1200.5
        assert run["trace_id"] == "trace-qa-001"
        # 非流式 + 自动路由(无 capability、无目标 → 数据集问数 + 智能选表)
        run2 = data(await c.post("/api/admin/quickbi/run", json={"question": "今天客流多少", "stream": False}))
        assert run2["capability"] == "dataset_qa" and run2["status"] == "success"
        # 入参错误以 HTTP 400 返回(不是塞进 200 的事件流)
        bad = await c.post("/api/admin/quickbi/run", json={"question": "  ", "stream": True})
        assert bad.status_code == 400


@run_async
async def test_dashboard_asset_and_fixed_questions():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD}))
        assert asset["ref_id"] == mock.PAGE_ID and asset["name"] == "经营驾驶舱"
        assert asset["cubes"][0]["cube_id"] == mock.CUBE_SALES
        # 候选问题一键转固定问题
        adopted = data(await c.post(f"/api/admin/quickbi/assets/{asset['id']}/adopt-questions", json={}))["items"]
        assert {q["capability"] for q in adopted} == {"dataset_qa", "insight"}
        # 手工新建带变量的固定问题(小Q报表 · 仪表板问数)
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "月度平台销售", "question": "{{month}} 各平台销售额",
            "capability": "dashboard_qa", "target": {"asset_id": asset["id"]}, "category": "经营"}))
        miss = await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run", json={"stream": False})
        assert miss.status_code == 400 and "month" in miss.json()["message"]
        run = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run",
                                json={"params": {"month": "2026年8月"}, "stream": False}))
        assert run["status"] == "success" and run["question"] == "2026年8月 各平台销售额"
        assert run["mode"] == "fixed" and run["fixed_question_id"] == fq["id"]
        # 仪表板资产缓存了数据集 → 不再重复拉取仪表板结构
        assert mock.STATE.calls.count(("GET", "/openapi/v2/dashboard/query")) == 1
        # 批量执行:解读(仪表板快照)+ 问数
        ids = [q["id"] for q in adopted]
        results = data(await c.post("/api/admin/quickbi/fixed-questions/batch-run", json={"ids": ids}))["items"]
        assert [r["status"] for r in results] == ["success", "success"], results
        insight = next(r for r in results if "环比" in (r.get("answer") or ""))
        assert insight
        fqs = data(await c.get("/api/admin/quickbi/fixed-questions", params={"category": "经营"}))["items"]
        assert fqs[0]["last_run_status"] == "success"


@run_async
async def test_chat_session_sticky_target_and_routing():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        f = await c.post("/api/admin/quickbi/files", data={"connection_id": conn["id"]},
                         files={"file": ("销售.csv", "月份,销售额\n1月,100\n2月,150\n".encode(), "text/csv")})
        file = data(f)
        assert file["qa_ready"] and file["insight_ready"] and file["capabilities"] == ["file_qa", "insight", "report", "doc_parse"]
        s = data(await c.post("/api/admin/quickbi/sessions", json={"target": {"file_id": file["id"]}}))
        r1 = data(await c.post(f"/api/admin/quickbi/sessions/{s['id']}/messages",
                               json={"question": "各部门人数TOP", "stream": False}))
        assert r1["capability"] == "file_qa" and r1["status"] == "success"
        r2 = data(await c.post(f"/api/admin/quickbi/sessions/{s['id']}/messages",
                               json={"question": "解读一下趋势", "stream": False}))
        assert r2["capability"] == "insight" and "环比" in r2["answer"]
        # 换成数据集会话:第一轮智能选表,选中的数据集粘在会话上,追问不再选表
        s2 = data(await c.post("/api/admin/quickbi/sessions", json={"capability": "dataset_qa"}))
        await c.post(f"/api/admin/quickbi/sessions/{s2['id']}/messages", json={"question": "上周客流", "stream": False})
        before = mock.STATE.calls.count(("POST", "/openapi/v2/smartq/tableSearch"))
        r4 = data(await c.post(f"/api/admin/quickbi/sessions/{s2['id']}/messages",
                               json={"question": "那销售额呢", "stream": False}))
        assert mock.STATE.calls.count(("POST", "/openapi/v2/smartq/tableSearch")) == before
        assert r4["target"]["cube_id"] == mock.CUBE_TRAFFIC
        assert mock.STATE.last_body["userQuestion"] == "上周客流;追问:那销售额呢"   # 追问带上文
        detail = data(await c.get(f"/api/admin/quickbi/sessions/{s2['id']}"))
        assert detail["turn_count"] == 2 and len(detail["messages"]) == 2 and detail["title"] == "上周客流"


@run_async
async def test_report_async_refresh_and_stream():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        f = data(await c.post("/api/admin/quickbi/files",
                              files={"file": ("经营数据.csv", "月份,收入\n1月,10\n".encode(), "text/csv")}))
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": mock.PAGE_ID}))
        started = data(await c.post("/api/admin/quickbi/reports", json={
            "connection_id": conn["id"], "question": "生成8月经营分析报告",
            "target": {"file_ids": [f["id"]], "asset_id": asset["id"]}}))
        assert started["status"] == "running" and started["chat_id"]
        res = mock.STATE.last_body["resources"]
        # 引用 = 上传文件 + 仪表板快照转出的 CSV
        assert res[0] == {"id": "ref-经营数据", "title": "经营数据", "type": "csv"} and len(res) == 2
        polls = mock.STATE.report_polls
        r1 = data(await c.get(f"/api/admin/quickbi/reports/{started['id']}"))
        assert r1["status"] == "running" and mock.STATE.report_polls == polls + 1
        # SSE 跟踪到结束
        s = await c.get(f"/api/admin/quickbi/reports/{started['id']}/stream")
        evs = sse_events(s)
        done = evs[-1]
        assert done["type"] == "done" and done["status"] == "success"
        assert done["report_url"] == f"https://bi.aliyun.com/copilot/qreportReplay?caseId={started['chat_id']}"
        final = data(await c.get(f"/api/admin/quickbi/reports/{started['id']}"))
        assert final["status"] == "success" and final["report_url"] == done["report_url"]
        assert final["answer"] == "报告正文"


@run_async
async def test_open_api_with_api_key():
    from cogni_agent.gateway_service.api_key_service import create_api_key

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "平台销售", "question": "各平台销售额",
            "capability": "dataset_qa", "target": {"cube_id": mock.CUBE_SALES}, "visibility": "public"})
        key = (await create_api_key(scenic_id=None, created_by="u-admin"))["key"]
        h = {"Authorization": f"Bearer {key}"}
        assert (await c.get("/v1/quickbi/connections")).status_code == 401
        conns = data(await c.get("/v1/quickbi/connections", headers=h))["items"]
        assert conns[0]["id"] == conn["id"] and "api_secret" not in conns[0]
        q = data(await c.post("/v1/quickbi/query", headers=h, json={
            "question": "各平台销售额", "target": {"cube_id": mock.CUBE_SALES}}))
        assert q["status"] == "success" and q["source"] == "api" and q["result"]["charts"]
        fqs = data(await c.get("/v1/quickbi/fixed-questions", headers=h))["items"]
        run = data(await c.post(f"/v1/quickbi/fixed-questions/{fqs[0]['id']}/run", headers=h, json={}))
        assert run["status"] == "success"
        sess = data(await c.post("/v1/quickbi/sessions", headers=h, json={"target": {"cube_id": mock.CUBE_SALES}}))
        stream = await c.post(f"/v1/quickbi/sessions/{sess['id']}/messages", headers=h,
                              json={"question": "各平台销售额", "stream": True})
        assert sse_events(stream)[-1]["status"] == "success"
        rep = data(await c.post("/v1/quickbi/reports", headers=h, json={"question": "生成销售报告"}))
        assert rep["status"] == "running"
        for _ in range(3):
            got = data(await c.get(f"/v1/quickbi/reports/{rep['id']}", headers=h))
        assert got["status"] == "success" and got["report_url"]
        # 关闭 API 开放后不可调用
        await c.put(f"/api/admin/quickbi/connections/{conn['id']}", json={"api_enabled": False})
        denied = await c.post("/v1/quickbi/query", headers=h, json={"question": "x", "connection_id": conn["id"]})
        assert denied.status_code == 404
        # 其他 Key 看不到别人的运行记录
        other = (await create_api_key(scenic_id=None, created_by="u-admin"))["key"]
        assert (await c.get(f"/v1/quickbi/runs/{q['id']}", headers={"Authorization": f"Bearer {other}"})
                ).status_code == 404


@run_async
async def test_permissions_between_users():
    viewer = CurrentUser(user_id="u-viewer", scenic_id=None, role="group_viewer")
    scenic = CurrentUser(user_id="u-scenic", scenic_id="S001", role="scenic_operator")
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
    async with _client(_app(viewer)) as c:
        assert [x["id"] for x in data(await c.get("/api/admin/quickbi/connections"))["items"]] == [conn["id"]]
        r = await c.put(f"/api/admin/quickbi/connections/{conn['id']}", json={"name": "x"})
        assert r.status_code == 403
    async with _client(_app(scenic)) as c:
        # 集团级(未绑景区)且 group 可见的连接 = 所有景区共用的那条默认连接,景区账号能用。
        # 这一条是按「景区绑定」的口径改过来的:现场只有一套 Quick BI 凭证、连接都挂在集团级,
        # 若对景区账号不可见,景区用户一发起就失败(而报错是「连接不存在」,没人会往可见范围上想)。
        assert [x["id"] for x in data(await c.get("/api/admin/quickbi/connections"))["items"]] == [conn["id"]]
        # 能用不等于能改:改连接仍然只有创建者/集团管理员可以
        assert (await c.put(f"/api/admin/quickbi/connections/{conn['id']}", json={"name": "x"})).status_code == 403


@run_async
async def test_connection_bound_to_other_scenic_stays_invisible():
    """绑到别的景区的连接,本景区账号看不到 —— 景区隔离不能因为上一条放宽而失效。"""
    from cogni_agent.common.db import get_sessionmaker
    from cogni_agent.tool_center_service.quickbi.models import QbiConnection

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
    async with get_sessionmaker()() as db:
        row = await db.get(QbiConnection, conn["id"])
        row.scenic_id = "S002"          # 归属别的景区
        await db.commit()

    scenic = CurrentUser(user_id="u-scenic", scenic_id="S001", role="scenic_operator")
    async with _client(_app(scenic)) as c:
        assert data(await c.get("/api/admin/quickbi/connections"))["items"] == []
        r = await c.post("/api/admin/quickbi/run", json={"connection_id": conn["id"], "question": "x"})
        assert r.status_code == 404


@run_async
async def test_publish_quickbi_tool_and_execute_through_engine():
    from cogni_agent.runtime_service.engine.resolver import resolve_execution_spec
    from cogni_agent.tool_center_service.mcp.protocol import invoke_tool
    from cogni_agent.tool_center_service.tool.templates import all_templates

    assert any(t["type"] == "quickbi" for t in all_templates()["templates"])
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD}))
        tool = data(await c.post("/api/admin/quickbi/tools", json={
            "name": "qbi_dashboard_qa", "connection_id": conn["id"], "capability": "dashboard_qa",
            "target": {"asset_id": asset["id"]}}))
        assert tool["type"] == "quickbi"
    async with common_db.get_sessionmaker()() as db:
        spec = await resolve_execution_spec(db, executor_type="tool", executor_id=tool["id"], scenic_id=None,
                                            role="super_admin", user_id="u-admin")
        ts = spec.direct_tool
        assert ts.type == "quickbi"
        assert ts.json_schema["required"] == ["question"]
        assert "dashboard" not in ts.json_schema["properties"]      # 目标已固定,不再让模型填
        out = await invoke_tool(db, ts.id, "", "super_admin", {"question": "本月各平台销售额"},
                                user_id="u-admin", type_config=ts.type_config, tool_type=ts.type)
    assert out["ok"] and out["status"] == "success", out
    assert "| 销售平台 | 销售额(SUM) |" in out["content"] and "| 天猫 | 1200.5 |" in out["content"]
    async with _client(_app()) as c:
        runs = data(await c.get("/api/admin/quickbi/runs", params={"mode": "tool"}))["items"]
        assert runs and runs[0]["source"] == "agent"

    # 所有者降级为景区账号、且连接被改绑到**别的景区**之后,工具引用的连接对所有者不再可见
    # → 执行时拒绝(不因调用者是超管而放行)。
    # 这里改的是归属景区而不只是 owner:集团级(未绑景区)的 group 连接现在是所有景区共用的,
    # 只换 owner 并不会让它对景区账号消失 —— 那正是「景区绑定」的口径。
    from cogni_agent.user_center_service.models import User
    async with common_db.get_sessionmaker()() as db:
        owner = await db.get(User, ADMIN.user_id)
        owner.role, owner.scenic_id = "scenic_operator", "S001"
        row = await db.get(QbiConnection, conn["id"])
        row.owner_user_id, row.scenic_id = "u-other", "S002"
        await db.commit()
        out2 = await invoke_tool(db, ts.id, "", "super_admin", {"question": "本月各平台销售额"},
                                 user_id="u-admin", type_config=ts.type_config, tool_type=ts.type)
    assert out2["ok"] is False and "所有者" in out2["error"]


@run_async
async def test_generic_tool_route_rejects_foreign_connection():
    from cogni_agent.common.db import get_sessionmaker
    from cogni_agent.tool_center_service.routes import ToolCreateRequest, create_tool

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
    # 把连接绑到别的景区,再用 S001 的账号去引用它 —— 这才是「借用别人的连接」
    from cogni_agent.tool_center_service.quickbi.models import QbiConnection

    async with get_sessionmaker()() as db:
        row = await db.get(QbiConnection, conn["id"])
        row.scenic_id = "S002"
        await db.commit()

    scenic = CurrentUser(user_id="u-scenic", scenic_id="S001", role="scenic_operator")
    req = ToolCreateRequest(name="steal", type="quickbi", type_config={"connection_id": conn["id"]})
    async with get_sessionmaker()() as db:
        with pytest.raises(ZhilvError):
            await create_tool(req, user=scenic, db=db)


@run_async
async def test_client_disconnect_finalizes_run():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
    actor = svc.Actor(user_id="u-admin", scenic_id=None, role="super_admin")
    agen = svc.run_stream(actor, connection_id=conn["id"], capability="dataset_qa", question="各平台销售额",
                          target={"cube_id": mock.CUBE_SALES})
    first = await agen.__anext__()
    await agen.__anext__()
    await agen.aclose()                      # 模拟 SSE 客户端中途断开
    run = await svc.get_run(actor, first["run_id"])
    assert run["status"] == "failed" and run["error_code"] == "cancelled"


@run_async
async def test_auto_routed_report_does_not_block_non_stream():
    async with _client(_app()) as c:
        await _setup_connection(c)
        run = data(await c.post("/api/admin/quickbi/run", json={"question": "帮我生成一份经营分析报告", "stream": False}))
        assert run["capability"] == "report" and run["status"] == "running" and run["chat_id"]
        assert mock.STATE.report_polls == 0          # 创建即返回,不等待进度


@run_async
async def test_server_domain_whitelist():
    async with _client(_app()) as c:
        r = await c.post("/api/admin/quickbi/connections", json={"name": "x", "server_domain": "http://10.0.0.8:8080"})
        assert r.status_code == 400 and "白名单" in r.json()["message"]
        ok = await c.post("/api/admin/quickbi/connections", json={"name": "y", "server_domain": "https://quickbi-public.cn-hangzhou.aliyuncs.com"})
        assert ok.status_code == 200


@run_async
async def test_env_default_connection(monkeypatch):
    s = get_settings().raw["quickbi"]
    s["default_connection"] = {"user_token": mock.USER_TOKEN, "server_domain": "https://bi.aliyun.com"}
    async with _client(_app()) as c:
        items = data(await c.get("/api/admin/quickbi/connections"))["items"]
        assert items[-1]["id"] == svc.ENV_CONNECTION_ID and items[-1]["virtual"] and items[-1]["trial_active"]
        run = data(await c.post("/api/admin/quickbi/run", json={
            "question": "各平台销售额", "target": {"cube_id": mock.CUBE_SALES}, "stream": False}))
        assert run["connection_id"] == svc.ENV_CONNECTION_ID and run["status"] == "success"
        r = await c.put(f"/api/admin/quickbi/connections/{svc.ENV_CONNECTION_ID}", json={"name": "x"})
        assert r.status_code in (403, 404)


def test_sql_sanity():
    # 迁移与模型的一致性由 scripts/init_db.py --check / 启动体检保证(tests/test_legacy_db_convergence.py
    # 另有真实库上的收敛回归);这里只确认 Quick BI 那条迁移已登记且编号唯一
    from cogni_agent.common.schema import MIGRATIONS
    ids = [m.id for m in MIGRATIONS]
    assert "0066_quickbi_integration" in ids and len(ids) == len(set(ids))
    qbi = next(m for m in MIGRATIONS if m.id == "0066_quickbi_integration")
    assert "tc_qbi_runs" in text(qbi.sql).text


@run_async
async def test_env_connection_hidden_from_scenic_users():
    get_settings().raw["quickbi"]["default_connection"] = {"user_token": mock.USER_TOKEN}
    scenic = CurrentUser(user_id="u-scenic", scenic_id="S001", role="scenic_operator")
    async with _client(_app(scenic)) as c:
        assert data(await c.get("/api/admin/quickbi/connections"))["items"] == []
        r = await c.post("/api/admin/quickbi/run", json={"connection_id": svc.ENV_CONNECTION_ID, "question": "x"})
        assert r.status_code == 403


# ─────────── 景区绑定 / 固定问题直达 / 思考与数据分离 ───────────

@run_async
async def test_asset_bound_to_scenic_and_not_auto_picked_on_ask():
    """资产登记时绑景区、按景区筛;提问只给 scenic_id **不会**替用户挂一块看板。

    以前会按景区兜底挑一个已登记资产,而「给没给目标数据」只认单数 cube_id ——
    多选的 cube_ids 被当成没给,挂上的看板把运营选的数据整个换掉。没给就是没给,
    由能力层报「缺数据来源」。
    """
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        # 一块绑到景区 wys 的仪表板 + 一块集团级的
        bound = data(await c.post("/api/admin/quickbi/assets/dashboards",
                                  json={"url": TOKEN3RD, "name": "武夷山驾驶舱", "scenic_id": "wys"}))
        assert bound["scenic_id"] == "wys"
        # 按景区筛:本景区 + 集团级都在,别的景区的不在
        only_wys = data(await c.get("/api/admin/quickbi/assets", params={"scenic_id": "wys"}))["items"]
        assert bound["id"] in [a["id"] for a in only_wys]
        other = data(await c.get("/api/admin/quickbi/assets", params={"scenic_id": "other"}))["items"]
        assert bound["id"] not in [a["id"] for a in other]

        # 提问不给 target,只给景区 → run 记录落到该景区,但不替用户挑看板
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dashboard_qa",
            "question": "本月客流环比", "scenic_id": "wys", "stream": False}))
        assert run["status"] == "failed" and "仪表板" in (run["error"] or "")
        assert run["scenic_id"] == "wys"
        assert not run["target"].get("asset_id")


@run_async
async def test_scenic_user_cannot_bind_or_ask_for_another_scenic():
    """景区角色越界要被挡:资产按景区找,能随手绑到别人景区就等于跨景区串数据。"""
    scenic_user = CurrentUser(user_id="u-wys", scenic_id="wys", role="scenic_admin")
    async with _client(_app(scenic_user)) as c:
        await _setup_connection(c)
        bad = await c.post("/api/admin/quickbi/assets/dashboards",
                           json={"url": TOKEN3RD, "scenic_id": "other"})
        assert bad.status_code == 403, bad.text
        bad_run = await c.post("/api/admin/quickbi/run", json={
            "capability": "dashboard_qa", "question": "本月客流", "scenic_id": "other", "stream": False})
        assert bad_run.status_code == 403, bad_run.text


@run_async
async def test_run_by_fixed_question_id_only():
    """只给固定问题 id 就能问:问题文本 / 能力 / 目标数据都从那条配置带出来。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD}))
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "本月客流", "question": "本月各平台销售额",
            "capability": "dashboard_qa", "target": {"asset_id": asset["id"]}}))
        run = data(await c.post("/api/admin/quickbi/run",
                                json={"fixed_question_id": fq["id"], "stream": False}))
        assert run["status"] == "success"
        assert run["question"] == "本月各平台销售额"
        assert run["capability"] == "dashboard_qa"
        assert run["fixed_question_id"] == fq["id"]
        assert run["target"].get("asset_id") == asset["id"]


@run_async
async def test_fixed_question_fields_can_be_overridden_per_call():
    """固定问题只是默认值:这一问显式带了能力/目标数据,就以调用方传的为准。

    操练台的能力条与目标数据下拉在选中固定问题后仍然可改 —— 改了要真的生效,
    否则那几个控件就是摆设。
    """
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD}))
        ds = data(await c.post("/api/admin/quickbi/assets/datasets",
                               json={"cube_id": "cube-sales-0001", "name": "订单销售明细"}))
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "本月客流", "question": "本月各平台销售额",
            "capability": "dashboard_qa", "target": {"asset_id": asset["id"]}}))
        run = data(await c.post("/api/admin/quickbi/run", json={
            "fixed_question_id": fq["id"], "capability": "dataset_qa",
            "target": {"asset_id": ds["id"], "asset_kind": "dataset"}, "stream": False}))
        assert run["status"] == "success"
        assert run["question"] == "本月各平台销售额"       # 问题文本仍来自固定问题
        assert run["capability"] == "dataset_qa"          # 能力以调用方为准
        assert run["target"].get("asset_id") == ds["id"]  # 目标数据同理
        assert run["fixed_question_id"] == fq["id"]


@run_async
async def test_events_carry_channel_and_result_splits_thinking_from_data():
    """每个事件带 channel,结果里 thinking / data 分开 —— 前端据此分区渲染。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        r = await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dataset_qa", "question": "各平台销售额",
            "target": {"cube_id": mock.CUBE_SALES}, "stream": True})
        evs = sse_events(r)
        typed = [e for e in evs if e["type"] not in ("run", "done")]
        assert typed, evs
        assert all(e.get("channel") in ("thinking", "data") for e in typed), typed
        # 结论正文(answer/conclusion/summary/report)算数据,推理/SQL 算过程
        assert any(e["type"] == "text" and e["channel"] == "data"
                   and e.get("section", "answer") in ("answer", "conclusion", "summary", "report")
                   for e in typed)
        assert any(e["type"] == "chart" and e["channel"] == "data" for e in typed)
        assert all(e["channel"] == "thinking" for e in typed if e["type"] in ("reasoning", "sql", "step", "status"))

        run_id = next(e["run_id"] for e in evs if e["type"] == "run")
        detail = data(await c.get(f"/api/admin/quickbi/runs/{run_id}"))
        res = detail["result"]
        assert res["data"]["answer"] == detail["answer"]
        assert "charts" in res["data"] and "reasoning" in res["thinking"]
        # 过程不该混进数据区
        assert "reasoning" not in res["data"] and "answer" not in res["thinking"]


# ─────────── 数据域 / 请求类别 / 对外目录 ───────────

@run_async
async def test_fixed_question_domain_and_ask_type():
    """数据域是自由文本(现场可新增),请求类别不填时按能力推导。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        a = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "本月票务收入", "question": "本月各渠道门票收入",
            "capability": "dashboard_qa", "domain": "票务"}))
        assert a["domain"] == "票务"
        # 仪表板问数 = 报表,不是问数:对调用方来说这是两类不同的东西
        assert a["ask_type"] == "dashboard" and a["ask_type_label"] == "报表"
        d = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "数据集问数", "question": "各平台销售额",
            "capability": "dataset_qa", "domain": "OTA"}))
        assert d["ask_type"] == "qa" and d["ask_type_label"] == "问数"
        assert d["domain"] == "OTA"

        b = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "经营解读", "question": "解读核心指标变化",
            "capability": "insight", "domain": "客群"}))
        assert b["ask_type"] == "insight" and b["ask_type_label"] == "解读"

        # 预置之外的域照样能存 —— 业务分类总会长出新的一类,不该要求改代码
        c3 = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "自定义域", "question": "随便问问",
            "capability": "dataset_qa", "domain": "研学"}))
        assert c3["domain"] == "研学"

        # 候选清单 = 预置 + 现场用过的
        doms = data(await c.get("/api/admin/quickbi/fixed-questions/domains"))
        assert "票务" in doms["presets"] and "画像" in doms["presets"] and "OTA" in doms["presets"]
        assert "研学" in doms["items"], doms["items"]
        assert {t["key"] for t in doms["ask_types"]} == {"qa", "dashboard", "insight", "report"}

        # 按域 / 按类别筛
        only_ticket = data(await c.get("/api/admin/quickbi/fixed-questions", params={"domain": "票务"}))["items"]
        assert [q["id"] for q in only_ticket] == [a["id"]]
        only_board = data(await c.get("/api/admin/quickbi/fixed-questions",
                                      params={"ask_type": "dashboard"}))["items"]
        assert [q["id"] for q in only_board] == [a["id"]]
        only_insight = data(await c.get("/api/admin/quickbi/fixed-questions", params={"ask_type": "insight"}))["items"]
        assert [q["id"] for q in only_insight] == [b["id"]]

        # 显式改类别:同一条问数问题也可以标成解读
        upd = data(await c.put(f"/api/admin/quickbi/fixed-questions/{a['id']}", json={"ask_type": "insight"}))
        assert upd["ask_type"] == "insight"
        # 清空数据域
        cleared = data(await c.put(f"/api/admin/quickbi/fixed-questions/{a['id']}", json={"domain": ""}))
        assert cleared["domain"] is None


@run_async
async def test_catalog_returns_questions_with_usable_data_ids():
    """对外目录:每条问题都带着「可以直接拿去提问的数据标识」,而不只是个名字。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards",
                                  json={"url": TOKEN3RD, "name": "经营驾驶舱", "scenic_id": "wys"}))
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "本月票务收入", "question": "本月各渠道门票收入",
            "capability": "dashboard_qa", "domain": "票务", "target": {"asset_id": asset["id"]}}))

        cat = data(await c.get("/api/admin/quickbi/catalog",
                               params={"scenic_id": "wys", "ask_type": "dashboard"}))
        assert cat["scenic_id"] == "wys"
        assert {t["key"] for t in cat["ask_types"]} == {"qa", "dashboard", "insight", "report"}
        q = next(x for x in cat["questions"] if x["id"] == fq["id"])
        assert q["domain"] == "票务" and q["ask_type"] == "dashboard"
        # 关键:page_id 能直接用于下一步提问
        assert q["target"]["kind"] == "dashboard"
        assert q["target"]["page_id"] == mock.PAGE_ID
        assert q["target"]["asset_id"] == asset["id"]
        # 该景区的全部仪表板/数据集也一并给出
        assert any(d["page_id"] == mock.PAGE_ID for d in cat["dashboards"])

        # 用目录里的 page_id 自己组问题,应当照常执行
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dashboard_qa", "question": "本月客流环比",
            "scenic_id": "wys", "target": {"page_id": q["target"]["page_id"]}, "stream": False}))
        assert run["status"] == "success"

        # 按类别筛:解读类的问题一条都没有,questions 应为空但结构完整
        empty = data(await c.get("/api/admin/quickbi/catalog", params={"scenic_id": "wys", "ask_type": "insight"}))
        assert empty["questions"] == [] and "dashboards" in empty


@run_async
async def test_catalog_open_api_shape():
    """开放平台的 /catalog 与管理端同一份数据(景区 Key 只看得到自己景区)。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD, "scenic_id": "wys"})
        await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "本月票务收入", "question": "本月各渠道门票收入",
            "capability": "dashboard_qa", "domain": "票务", "visibility": "public"})
    async with _client(_app()) as c:
        r = await c.get("/v1/quickbi/catalog", params={"scenic_id": "wys"},
                        headers={"Authorization": "Bearer sk-zhilv-test"})
        # 没有真实 API Key 时应当是 401,而不是 500 —— 鉴权失败要能明确区分
        assert r.status_code in (401, 403), r.text


@run_async
async def test_thinking_prefix_is_split_out_of_the_answer():
    """小Q把思维链和回答塞在同一段文本里时,回答气泡里不该再出现「思考过程」。"""
    from cogni_agent.tool_center_service.quickbi.events import AnswerSplitter

    blob = ("思考过程: 1. 第一步:用户意图识别\n2. 第二步:按类型C输出\n"
            "完美,开始输出。### 用户问题解答\n针对您提出的问题,直接解答如下:\n- 核心指标变化:住宿消费升级\n")
    sp = AnswerSplitter()
    out = []
    for ch in blob:                      # 逐字符:最接近真实流式,也最容易把分界标志切碎
        out += sp.feed({"type": "text", "delta": ch, "section": "answer"})
    out += sp.flush()
    reasoning = "".join(e["delta"] for e in out if e["type"] == "reasoning")
    answer = "".join(e["delta"] for e in out if e["type"] == "text")
    assert reasoning.startswith("思考过程") and "第二步" in reasoning
    assert answer.startswith("### 用户问题解答")
    assert "思考过程" not in answer and "第二步" not in answer


# ─────────── 报告后台续跑 ───────────

@run_async
async def test_report_watcher_finishes_report_without_any_client():
    """客户端断开之后,报告结果仍然要由服务端自己收回来。

    这是「报告要很久、必须保证一直获取结果」那条需求的核心回归:
    发起报告 → 谁都不去跟踪(模拟页面关掉)→ 后台扫一轮 → 记录应当已收尾。
    """
    from cogni_agent.tool_center_service.quickbi import report_watcher

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": mock.PAGE_ID}))
        started = data(await c.post("/api/admin/quickbi/reports", json={
            "connection_id": conn["id"], "question": "生成8月经营分析报告",
            "target": {"asset_id": asset["id"]}}))
        assert started["status"] == "running" and started["chat_id"]

        # 没有任何人在跟踪这条报告 —— 后台扫描是唯一的收尾途径。
        # 一轮扫描 = 拉一次进度,远端没生成完就下一轮再拉(常驻循环每 30s 一轮)。
        finished = 0
        for _ in range(10):
            stats = await report_watcher.sweep_once()
            assert stats["errors"] == 0, stats
            finished += stats["finished"]
            if finished:
                break
        assert finished >= 1, "后台扫描没能把报告收尾"

        final = data(await c.get(f"/api/admin/quickbi/runs/{started['id']}"))
        assert final["status"] == "success", final
        assert final["report_url"].endswith(started["chat_id"])
        assert final["answer"] == "报告正文"

        # 幂等:已完成的记录不再进入待办,再扫也不会重复收尾
        again = await report_watcher.sweep_once()
        assert again["finished"] == 0, again


@run_async
async def test_report_watcher_gives_up_with_a_clear_reason():
    """超过兜底时限仍无结论的,要给明确结论 —— 永远 running 的记录比失败更难查。"""
    from datetime import timedelta

    from cogni_agent.common.db import get_sessionmaker, utcnow
    from cogni_agent.tool_center_service.quickbi import report_watcher
    from cogni_agent.tool_center_service.quickbi.models import QbiRun

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        started = data(await c.post("/api/admin/quickbi/reports", json={
            "connection_id": conn["id"], "question": "一份很久以前发起的报告",
            "target": {"data": "月份,收入\n1月,10\n"}}))
        # 把创建时间推到兜底时限之外
        async with get_sessionmaker()() as db:
            row = await db.get(QbiRun, started["id"])
            row.created_at = utcnow() - timedelta(hours=48)
            await db.commit()

        stats = await report_watcher.sweep_once()
        assert stats["gave_up"] >= 1, stats

        final = data(await c.get(f"/api/admin/quickbi/runs/{started['id']}"))
        assert final["status"] == "failed"
        assert final["error_code"] == "REPORT_GIVE_UP"
        assert started["chat_id"] in final["error"]          # 给出 chatId,能去 Quick BI 侧自查
        assert "小时" in final["error"]


@run_async
async def test_report_watcher_survives_a_broken_record():
    """一条坏记录不能拖垮整轮扫描 —— 其余报告照样要收。"""
    from cogni_agent.common.db import get_sessionmaker
    from cogni_agent.tool_center_service.quickbi import report_watcher
    from cogni_agent.tool_center_service.quickbi.models import QbiRun

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        broken = data(await c.post("/api/admin/quickbi/reports", json={
            "connection_id": conn["id"], "question": "坏记录",
            "target": {"data": "月份,收入\n1月,10\n"}}))
        async with get_sessionmaker()() as db:            # 连接指向一个不存在的 id
            row = await db.get(QbiRun, broken["id"])
            row.connection_id = "conn-does-not-exist"
            await db.commit()
        good = data(await c.post("/api/admin/quickbi/reports", json={
            "connection_id": conn["id"], "question": "正常报告",
            "target": {"data": "月份,收入\n1月,10\n"}}))

        for _ in range(10):
            await report_watcher.sweep_once()
            if data(await c.get(f"/api/admin/quickbi/runs/{good['id']}"))["status"] != "running":
                break
        ok = data(await c.get(f"/api/admin/quickbi/runs/{good['id']}"))
        assert ok["status"] == "success", ok               # 好的那条照常收尾

        # 坏记录要**收尾**,不是每轮记一次错然后重试。连接都没了,它永远不可能有进展 ——
        # 一直重试的话,这条记录会安静地躺在「生成中」里被每 30 秒捞出来一次,直到天荒地老。
        bad = data(await c.get(f"/api/admin/quickbi/runs/{broken['id']}"))
        assert bad["status"] == "failed", bad
        assert "已被删除" in (bad.get("error") or ""), bad
        # 收尾之后不再被捞:再扫一轮,checked 里不该还有它
        before = await report_watcher.sweep_once()
        assert broken["id"] not in str(before), before


@run_async
async def test_report_survives_polling_timeout_and_keeps_being_fetched():
    """拉进度超时 ≠ 报告失败。

    现场症状:`请求 Quick BI 超时: GET /openapi/v2/smartq/qreportChatData`。
    报告本身在远端好好跑着,只是这一次没拉回来。此时把运行记录判死,
    等于白白丢掉一份已经生成好的报告 —— 必须保持 running 并继续拉。
    """
    from cogni_agent.tool_center_service.quickbi import capabilities as cap
    from cogni_agent.tool_center_service.quickbi import report_watcher
    from cogni_agent.tool_center_service.quickbi.errors import QuickBIError

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        started = data(await c.post("/api/admin/quickbi/reports", json={
            "connection_id": conn["id"], "question": "生成经营分析报告",
            "target": {"data": "月份,收入\n1月,10\n"}}))
        assert started["status"] == "running"

        # 前两轮全部超时
        real = cap.fetch_report_raw
        calls = {"n": 0}

        async def flaky(qbi, chat_id, **kw):
            calls["n"] += 1
            if calls["n"] <= 2:
                raise QuickBIError("请求 Quick BI 超时:GET /openapi/v2/smartq/qreportChatData", kind="network")
            return await real(qbi, chat_id, **kw)

        cap.fetch_report_raw = flaky
        try:
            for _ in range(2):
                stats = await report_watcher.sweep_once()
                assert stats["errors"] >= 1 and stats["gave_up"] == 0, stats
                mid = data(await c.get(f"/api/admin/quickbi/runs/{started['id']}"))
                assert mid["status"] == "running", "超时把报告判死了"
                assert mid["error"] is None, mid["error"]
            # 网络恢复后照常收尾
            finished = 0
            for _ in range(10):
                finished += (await report_watcher.sweep_once())["finished"]
                if finished:
                    break
            assert finished >= 1
        finally:
            cap.fetch_report_raw = real

        final = data(await c.get(f"/api/admin/quickbi/runs/{started['id']}"))
        assert final["status"] == "success", final


@run_async
async def test_report_stops_only_on_a_real_quickbi_error():
    """Quick BI 明确报错(额度用尽这类)才停止跟踪 —— 一直重试没有意义。"""
    from cogni_agent.tool_center_service.quickbi import capabilities as cap
    from cogni_agent.tool_center_service.quickbi import report_watcher
    from cogni_agent.tool_center_service.quickbi.errors import QuickBIError

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        started = data(await c.post("/api/admin/quickbi/reports", json={
            "connection_id": conn["id"], "question": "生成经营分析报告",
            "target": {"data": "月份,收入\n1月,10\n"}}))

        real = cap.fetch_report_raw

        async def boom(qbi, chat_id, **kw):
            raise QuickBIError("小Q报告免费额度已用尽", code="AE0533330017", kind="business")

        cap.fetch_report_raw = boom
        try:
            stats = await report_watcher.sweep_once()
        finally:
            cap.fetch_report_raw = real
        assert stats["finished"] >= 1, stats

        final = data(await c.get(f"/api/admin/quickbi/runs/{started['id']}"))
        assert final["status"] == "failed"
        assert final["error_code"] == "AE0533330017"
        assert "额度" in final["error"]


@run_async
async def test_console_quickbi_turns_land_in_a_session():
    """操练台里问 Quick BI,这一问要能出现在会话列表,并且能按会话回放。

    前端的做法是:先建会话,再把每一轮打到这个会话下(固定问题走 /run 带 session_id,
    自由提问走 /sessions/{id}/messages)。这里用同样的调用序列锁住后端这一侧。
    """
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards",
                                  json={"url": TOKEN3RD, "scenic_id": "wys"}))
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "本月客流", "question": "本月各平台销售额",
            "capability": "dashboard_qa", "target": {"asset_id": asset["id"]}}))

        s = data(await c.post("/api/admin/quickbi/sessions", json={
            "connection_id": conn["id"], "capability": "dashboard_qa",
            "target": {"asset_id": asset["id"]}, "title": "本月客流", "scenic_id": "wys"}))
        assert s["scenic_id"] == "wys"

        # 第一轮:自由提问
        r1 = data(await c.post(f"/api/admin/quickbi/sessions/{s['id']}/messages",
                               json={"question": "本月各平台销售额", "stream": False}))
        assert r1["status"] == "success" and r1["session_id"] == s["id"]
        # 会话粘住了景区:这一轮的运行记录也归在该景区名下
        assert r1["scenic_id"] == "wys"

        # 第二轮:固定问题(走 /run,但归到同一个会话)
        r2 = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "fixed_question_id": fq["id"],
            "session_id": s["id"], "stream": False}))
        assert r2["session_id"] == s["id"] and r2["fixed_question_id"] == fq["id"]

        # 列表里能看到它,轮次数对得上
        listed = data(await c.get("/api/admin/quickbi/sessions"))["items"]
        mine = next(x for x in listed if x["id"] == s["id"])
        assert mine["turn_count"] >= 1

        # 打开会话能按时间顺序回放两轮
        detail = data(await c.get(f"/api/admin/quickbi/sessions/{s['id']}"))
        qs = [m["question"] for m in detail["messages"]]
        assert qs == ["本月各平台销售额", "本月各平台销售额"], qs
        assert [m["mode"] for m in detail["messages"]] == ["chat", "fixed"]

        # 归档后不再出现在列表里
        await c.delete(f"/api/admin/quickbi/sessions/{s['id']}")
        after = data(await c.get("/api/admin/quickbi/sessions"))["items"]
        assert all(x["id"] != s["id"] for x in after)


@run_async
async def test_nl2sql_prelude_goes_to_thinking_not_the_answer():
    """数据集问数的「怎么翻成 SQL」前言属于过程,不该出现在回答气泡里。

    现场形态(没有「思考过程」字样,但每行都是过程):
        问题选中数据集: … / 提及字段及维值: … / 判定0: … / 查询的逻辑SQL如下:
    """
    from cogni_agent.tool_center_service.quickbi.events import AnswerSplitter

    blob = ("问题选中数据集: 客群洞察-到访偏好(美食/购物/酒店/景点)\n"
            "提及字段及维值: 偏好类型(景点偏好);偏好对象;偏好占比;TGI值\n"
            "判定0: 按偏好对象看偏好占比与 TGI 值的分布。\n"
            "判定1: 单步 SQL;仅分组展示。\n"
            "查询的逻辑SQL如下:\n"
            "景点偏好中风景名胜占比最高(735.28),TGI 6144。\n")
    sp = AnswerSplitter()
    out = []
    for ch in blob:                       # 逐字符:分界标志被切碎也要认得出来
        out += sp.feed({"type": "text", "delta": ch, "section": "answer"})
    out += sp.flush()
    reasoning = "".join(e["delta"] for e in out if e["type"] == "reasoning")
    answer = "".join(e["delta"] for e in out if e["type"] == "text")
    assert "判定0" in reasoning and "问题选中数据集" in reasoning
    assert "风景名胜" in answer
    assert "判定" not in answer and "数据集" not in answer.split("\n")[0]
    # 通道标记也要对上:过程归 thinking,结论归 data
    from cogni_agent.tool_center_service.quickbi.events import event_channel
    assert event_channel({"type": "reasoning"}) == "thinking"
    assert event_channel({"type": "text", "section": "answer"}) == "data"


@run_async
async def test_nl2sql_prelude_alone_leaves_the_answer_empty_not_swallowed():
    """整段都是前言(结论在图表里)时:前言进过程,回答区空着,内容一个字都不能丢。"""
    from cogni_agent.tool_center_service.quickbi.events import AnswerSplitter

    blob = ("问题选中数据集: 客群洞察\n判定0: 分组展示两个度量。\n查询的逻辑SQL如下:\n")
    sp = AnswerSplitter()
    out = [e for ch in blob for e in sp.feed({"type": "text", "delta": ch, "section": "answer"})]
    out += sp.flush()
    reasoning = "".join(e["delta"] for e in out if e["type"] == "reasoning")
    answer = "".join(e["delta"] for e in out if e["type"] == "text")
    assert answer.strip() == "", answer
    # 丢没丢:去掉分界标志本身之后,原文应当原样出现在推理里
    assert "问题选中数据集" in reasoning and "判定0" in reasoning


# ─────────── 执行留痕与导出 ───────────

@run_async
async def test_run_debug_bundle_has_raw_events_and_no_secrets():
    """一次运行要能整包导出:上游原样事件 + 归一化事件 + 结果分区,且不带任何凭证。

    这是为了让「问完之后把执行过程发给研发」这件事成立 —— 光靠复述现象,
    最要命的一类问题(我们理解错了上游发的东西)根本看不出来。
    """
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dataset_qa", "question": "各平台销售额",
            "target": {"cube_id": mock.CUBE_SALES}, "stream": False}))

        bundle = data(await c.get(f"/api/admin/quickbi/runs/{run['id']}/debug"))
        assert bundle["captured"] is True
        assert bundle["run"]["question"] == "各平台销售额"
        assert bundle["run"]["capability_label"] == "问数"           # 四个字面,不带「小Q」
        assert bundle["run"]["capability_source"] == "数据集"
        # 上游原样事件确实留下来了,而且带着是哪个接口发的
        assert bundle["events"]["raw"], bundle["events"]
        assert any("smartq" in e["uri"] for e in bundle["events"]["raw"])
        # 归一化事件带通道标记,便于对照「我们是怎么理解的」
        assert any(e.get("channel") in ("thinking", "data") for e in bundle["events"]["normalized"])
        # 结果按两路分区
        assert "data" in bundle["result"] and "thinking" in bundle["result"]

        # 凭证绝不能出现在导出内容里 —— 这份东西是要发给别人的
        blob = json.dumps(bundle, ensure_ascii=False)
        for secret in (mock.USER_TOKEN, mock.API_SECRET):
            assert secret not in blob, "导出内容里泄漏了凭证"
        assert "***" in blob or "token" not in blob.lower()

        # Markdown 版本:给人读的
        md = await c.get(f"/api/admin/quickbi/runs/{run['id']}/debug", params={"format": "md"})
        assert md.status_code == 200 and "text/markdown" in md.headers["content-type"]
        text = md.text
        assert "# Quick BI 执行详情" in text and "各平台销售额" in text
        assert "上游原样事件" in text and "归一化事件" in text
        for secret in (mock.USER_TOKEN, mock.API_SECRET):
            assert secret not in text


@run_async
async def test_debug_trace_is_bounded_and_maskable():
    """留痕要限长、要脱敏:它是附加品,不能把库撑爆,更不能把凭证带出门。"""
    from cogni_agent.tool_center_service.quickbi.trace import MAX_EVENTS, RunTrace, mask

    # 脱敏:键名命中 / 文本里裸露的凭证都要打码
    masked = mask({"userToken": "deadbeefcafe4f0e8a1b2c3d4e5f6071",
                   "nested": {"api_secret": "5eaacf44", "ok": "keep"},
                   "text": "accessTicket=807cc156-9109-4dfc",
                   "list": [{"Authorization": "Bearer x"}]})
    blob = json.dumps(masked, ensure_ascii=False)
    assert "deadbeefcafe4f0e8a1b2c3d4e5f6071" not in blob
    assert "5eaacf44" not in blob and "807cc156" not in blob
    assert "keep" in blob                       # 不相干的字段照常保留

    # 限量:超过上限只记省略条数,不再堆内容(用不可合并的事件测上限)
    t = RunTrace()
    for i in range(MAX_EVENTS + 50):
        t.add_raw("/x", {"type": "status", "stage": f"s{i}"})
    d = t.to_dict()
    assert len(d["raw"]) == MAX_EVENTS and d["dropped"]["raw"] == 50

    # 合并:流式是逐字吐的,一句话能拆成上百条。同类型的连续增量并成一条 ——
    # 否则导出里一屏全是「用户」「要求解读」「提供的」,既撑爆限额也没人读得下去
    t2 = RunTrace()
    for piece in ("用户", "要求解读", "提供的数据"):
        t2.add_raw("/x", {"type": "summary", "subType": "model_reasoning", "data": piece})
        t2.add_normalized({"type": "reasoning", "delta": piece, "channel": "thinking"})
    t2.add_normalized({"type": "text", "delta": "结论", "section": "answer", "channel": "data"})
    d2 = t2.to_dict()
    assert len(d2["raw"]) == 1 and d2["raw"][0]["event"]["data"] == "用户要求解读提供的数据"
    assert d2["raw"][0]["_merged_deltas"] == 3
    # 推理与回答是两条,不会被并到一起
    assert [e["type"] for e in d2["normalized"]] == ["reasoning", "text"]
    assert d2["normalized"][0]["delta"] == "用户要求解读提供的数据"
    assert "_k" not in json.dumps(d2)          # 内部合并键不外泄

    # 关掉就一条不留
    off = RunTrace(enabled=False)
    off.add_raw("/x", {"a": 1})
    off.add_normalized({"type": "text"})
    assert off.to_dict() is None


@run_async
async def test_real_interpretation_stream_splits_reasoning_by_subtype():
    """回放现场导出的真实事件形态:推理必须进 thinking,回答才进 data。

    事件取自 2026-09-18 现场导出(quickbi-run-1ff7e362):
        {"type": "summary", "subType": "model_reasoning", "data": "用户要求解读"}
    推理与回答**共用 type**,只有 subType 能区分。以前不读这个字段,
    整条思维链被当成回答灌进气泡 —— 也就是用户看到的「思考过程和结果粘在一起」。
    """
    from cogni_agent.tool_center_service.quickbi.events import InterpretationNormalizer, with_channel

    upstream = [
        {"type": "trace", "data": "e2f1856d-efde-45ea-8bdc-6206247d61c2"},
        {"type": "locale", "data": "zh_CN"},
        {"type": "summary", "subType": "model_reasoning", "data": "用户"},
        {"type": "summary", "subType": "model_reasoning", "data": "要求解读"},
        {"type": "summary", "subType": "model_reasoning", "data": "提供的数据仪表板的核心指标"},
        {"type": "summary", "subType": "model_reasoning", "data": "。\n这属于“类型C:简单分析问题”。"},
        {"type": "summary", "subType": "model_reasoning", "data": "\n没有遗漏,逻辑清晰。直接输出。"},
        {"type": "text", "data": "【用户问题解答】\n"},
        {"type": "text", "data": "- **核心指标变化**:客群结构向“上班族”集中。"},
        {"type": "finish", "data": ""},
    ]
    norm = InterpretationNormalizer()
    out = []
    for e in upstream:
        out.extend(with_channel(x) for x in norm.feed(e))

    reasoning = "".join(e["delta"] for e in out if e["type"] == "reasoning")
    answer = "".join(e["delta"] for e in out if e["type"] == "text")

    assert "用户要求解读" in reasoning and "类型C" in reasoning and "直接输出" in reasoning
    assert answer.startswith("【用户问题解答】")
    assert "类型C" not in answer and "逻辑清晰" not in answer
    # 通道:推理 thinking、回答 data,前端据此分区
    assert all(e["channel"] == "thinking" for e in out if e["type"] == "reasoning")
    assert all(e["channel"] == "data" for e in out if e["type"] == "text")
    # locale 这类噪音不该产出任何事件
    assert not any(e.get("delta") == "zh_CN" for e in out)


def test_reasoning_subtype_detection_covers_upstream_spellings():
    """subType 的各种写法都要认得出来,别因为换个拼法又漏回回答区。"""
    from cogni_agent.tool_center_service.quickbi.events import is_reasoning_event

    for sub in ("model_reasoning", "MODEL_REASONING", "reasoning", "thinking",
                "chain_of_thought", "cot", "agent_reasoning", "model_thinking"):
        assert is_reasoning_event({"type": "summary", "subType": sub}), sub
    for sub in ("", None, "model_answer", "conclusion", "summary"):
        assert not is_reasoning_event({"type": "summary", "subType": sub}), sub


@run_async
async def test_debug_captures_raw_api_responses_and_masks_token_in_plain_fields():
    """留痕要覆盖「API 的第一版输出」,不只是 SSE。

    选表、仪表板结构、快照、报告进度都走非流式接口 —— 只留 SSE 的话,
    这些接口到底返回了什么无从查起,而问题恰恰常出在这一步。

    同时锁一个脱敏陷阱:Quick BI 把 user_token 放在 `userId` / `oapiUserId` 这种
    看不出是凭证的字段里,只靠键名匹配会漏 —— 必须按凭证字面值打码。
    """
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD}))
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "insight", "question": "趋势如何",
            "target": {"asset_id": asset["id"]}, "stream": False}))

        bundle = data(await c.get(f"/api/admin/quickbi/runs/{run['id']}/debug"))
        http = bundle["events"]["http"]
        assert http, "非流式接口一次都没留下来"
        # 仪表板快照那条链路的接口要在里面,且带着响应原文
        assert any("dashboard" in h["uri"] or "snapshot" in h["uri"] or "skills" in h["uri"] for h in http), \
            [h["uri"] for h in http]
        assert all("method" in h and "status" in h and "response" in h for h in http)

        # 凭证:既不能出现在 JSON 包里,也不能出现在 Markdown 里
        blob = json.dumps(bundle, ensure_ascii=False)
        assert mock.USER_TOKEN not in blob, "user_token 从 userId 这类字段漏出去了"
        assert mock.API_SECRET not in blob
        md = (await c.get(f"/api/admin/quickbi/runs/{run['id']}/debug", params={"format": "md"})).text
        assert mock.USER_TOKEN not in md and mock.API_SECRET not in md
        assert "接口原始一问一答" in md

        # 流式那一路仍在,而且推理与回答分开了
        assert bundle["events"]["raw"], "SSE 事件没留下来"
        assert (bundle["result"]["thinking"]["reasoning"] or "").strip(), "解读的思考过程没进 thinking"
        assert "判定为类型C" not in (bundle["result"]["data"]["answer"] or "")


@run_async
async def test_dashboard_qa_follows_official_dashboard_skill():
    """仪表板问数对齐官方「小Q报表技能」:
    问题点到某张图 → 用它的数据集;否则看板全部数据集走 multipleCubeIds 多表问数(不再先挑一份)。"""
    mock.STATE.two_cube_dashboard = True
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        asset = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD}))
        assert {x["cube_id"] for x in asset["cubes"]} == {mock.CUBE_SALES, mock.CUBE_TRAFFIC}

        # 1) 没点到具体图 → 多表问数,一次调用,两份数据都回来
        mock.STATE.calls.clear()
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dashboard_qa", "question": "整体经营怎么样",
            "target": {"asset_id": asset["id"]}, "stream": False}))
        assert run["status"] == "success", run
        paths = [p for _, p in mock.STATE.calls]
        assert paths.count("/openapi/v2/smartq/queryByQuestion") == 1
        assert "/openapi/v2/smartq/tableSearch" not in paths, "不再先智能挑一份"
        assert "/openapi/v2/snapshot/calling/shot" not in paths
        assert set(mock.STATE.last_body["multipleCubeIds"]) == {mock.CUBE_SALES, mock.CUBE_TRAFFIC}
        assert len(run["result"]["charts"]) == 2 and "客流平稳" in run["answer"]

        # 2) 问题点到了图表名 → 用这张图的数据集(单表流式)
        mock.STATE.calls.clear()
        run2 = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dashboard_qa", "question": "客流趋势最近怎么样",
            "target": {"asset_id": asset["id"]}, "stream": False}))
        assert run2["status"] == "success"
        assert mock.STATE.last_body.get("cubeId") == mock.CUBE_TRAFFIC
        assert "/openapi/v2/smartq/queryByQuestion" not in [p for _, p in mock.STATE.calls]

        # 3) 只选了销售那张图 → 候选只剩一个数据集,直接单表
        mock.STATE.calls.clear()
        run3 = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dashboard_qa", "question": "整体怎么样",
            "target": {"asset_id": asset["id"], "chart_ids": ["c1"]}, "stream": False}))
        assert run3["status"] == "success" and mock.STATE.last_body.get("cubeId") == mock.CUBE_SALES

        # 4) 显式「挑一份数据集」的存量配置:保留老跑法
        mock.STATE.calls.clear()
        run4 = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dashboard_qa", "question": "各平台销售额",
            "target": {"asset_id": asset["id"], "qa_scope": "dataset"}, "stream": False}))
        assert run4["status"] == "success"
        assert "/openapi/v2/smartq/tableSearch" in [p for _, p in mock.STATE.calls]


@run_async
async def test_multi_dataset_qa_uses_multiple_cube_ids():
    """固定问题 / 目标数据多选了数据集:一次 multipleCubeIds 多表问数。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        mock.STATE.calls.clear()
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "dataset_qa", "question": "销售和客流",
            "target": {"cube_ids": [mock.CUBE_SALES, mock.CUBE_TRAFFIC]}, "stream": False}))
        assert run["status"] == "success", run
        assert mock.STATE.last_body["multipleCubeIds"] == [mock.CUBE_SALES, mock.CUBE_TRAFFIC]  # 数组,不是逗号串(ec525170)
        assert len(run["result"]["charts"]) == 2
        # 解读 + 多选数据集:同样一次多表取数再解读
        run2 = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "insight", "question": "销售和客流解读",
            "target": {"cube_ids": [mock.CUBE_SALES, mock.CUBE_TRAFFIC]}, "stream": False}))
        assert run2["status"] == "success", run2


@run_async
async def test_dataset_asset_category():
    """数据集自定义分类:输入回车即存(不需要先建分类),可批量设,空串清掉。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        a = data(await c.post("/api/admin/quickbi/assets/datasets",
                              json={"connection_id": conn["id"], "cube_id": mock.CUBE_SALES}))
        b = data(await c.post("/api/admin/quickbi/assets/datasets",
                              json={"connection_id": conn["id"], "cube_id": mock.CUBE_TRAFFIC}))
        assert a.get("category") is None
        r = data(await c.post("/api/admin/quickbi/assets/category", json={"ids": [a["id"], b["id"]], "category": " 票务 "}))
        assert [x["category"] for x in r["items"]] == ["票务", "票务"]
        one = data(await c.put(f"/api/admin/quickbi/assets/{b['id']}", json={"category": "客流"}))
        assert one["category"] == "客流"
        lst = data(await c.get("/api/admin/quickbi/assets"))["items"]
        assert {x["id"]: x["category"] for x in lst} == {a["id"]: "票务", b["id"]: "客流"}
        cleared = data(await c.put(f"/api/admin/quickbi/assets/{b['id']}", json={"category": ""}))
        assert cleared["category"] is None


# ── 对齐 1.2.4 时修掉的几处目标数据 / 连接 / 变量问题 ─────────────────────

@run_async
async def test_multi_dataset_target_is_not_replaced_by_scenic_asset():
    """固定问题选了多份数据集(cube_ids),带着景区跑也必须问这几份 —— 不能被景区下登记的
    看板/数据集顶替(预热、APP 发起都会带景区)。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        data(await c.post("/api/admin/quickbi/assets/dashboards",
                          json={"url": TOKEN3RD, "name": "武夷山驾驶舱", "scenic_id": "wys"}))
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "多表", "question": "本月各平台销售额与客流",
            "capability": "dataset_qa", "target": {"cube_ids": [mock.CUBE_SALES, mock.CUBE_TRAFFIC]}}))
        mock.STATE.calls.clear()
        run = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run",
                                json={"scenic_id": "wys", "stream": False}))
        assert run["status"] == "success", run
        assert not run["target"].get("asset_id")
        assert set(mock.STATE.last_body["multipleCubeIds"]) == {mock.CUBE_SALES, mock.CUBE_TRAFFIC}
        # 列表上的「数据来源」按 target 实际给出,而不是能力的静态描述
        assert run["capability_source"] == "数据集 ×2"


@run_async
async def test_insight_prefers_selected_datasets_over_leftover_dashboard():
    """解读:target 里同时有残留的看板 asset_id 和刚勾的 cube_ids 时,按数据集先取数再解读,
    不去拍看板快照。"""
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        board = data(await c.post("/api/admin/quickbi/assets/dashboards", json={"url": TOKEN3RD}))
        mock.STATE.calls.clear()
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "insight", "question": "解读本月销售",
            "target": {"asset_id": board["id"], "cube_ids": [mock.CUBE_SALES, mock.CUBE_TRAFFIC]},
            "stream": False}))
        paths = [p for _, p in mock.STATE.calls]
        assert "/openapi/v2/snapshot/calling/shot" not in paths, paths
        assert run["status"] == "success", run


@run_async
async def test_usable_connection_never_borrows_another_scenics_connection():
    """问题绑的连接没了、本景区和集团级都没有可用连接时返回 None(报「去哪配」),
    而不是拿集团账号看得见的第一条 —— 那多半是别的景区的,会静默串数据。"""
    async with common_db.get_sessionmaker()() as db:
        db.add(QbiConnection(id="c-other", name="别的景区的连接", owner_user_id="u-admin",
                             visibility="group", scenic_id="other", active=True))
        await db.commit()
    group = svc.Actor(user_id="u-admin", scenic_id=None, role="super_admin", source="console")
    assert await svc.usable_connection_id(group, "gone", "wys") is None
    assert await svc.usable_connection_id(group, "gone", None) is None
    assert await svc.usable_connection_id(group, "gone", "other") == "c-other"


@run_async
async def test_runtime_vars_are_filled_not_reported_missing():
    """{{截止时间}} / {{统计日期}} / {{昨日}} 由运行时刻决定,调用方给不了 —— 不能算「缺参数」。"""
    from datetime import timedelta

    from cogni_agent.common.time import now_cst

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "昨日客流", "question": "{{昨日}} 客流,截至 {{截止时间}}",
            "capability": "dataset_qa", "target": {"cube_id": mock.CUBE_TRAFFIC}}))
        run = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run", json={"stream": False}))
        assert run["status"] == "success", run
        sent = mock.STATE.last_body["userQuestion"]
        assert "{{" not in sent
        assert (now_cst().date() - timedelta(days=1)).isoformat() in sent


@run_async
async def test_catalog_lists_multi_dataset_targets():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "多表", "question": "销售与客流",
            "capability": "dataset_qa", "target": {"cube_ids": [mock.CUBE_SALES, mock.CUBE_TRAFFIC]}}))
        cat = data(await c.get("/api/admin/quickbi/catalog"))
        q = next(x for x in cat["questions"] if x["id"] == fq["id"])
        assert q["target"]["kind"] == "dataset"
        assert q["target"]["cube_ids"] == [mock.CUBE_SALES, mock.CUBE_TRAFFIC]


@run_async
async def test_adopted_questions_follow_the_dashboards_scenic():
    """一键生成的固定问题跟着看板的景区走,不生成一批「集团级、全景区可见」却别人打不开的问题。"""
    async with _client(_app()) as c:
        await _setup_connection(c)
        board = data(await c.post("/api/admin/quickbi/assets/dashboards",
                                  json={"url": TOKEN3RD, "name": "武夷山驾驶舱", "scenic_id": "wys"}))
        created = data(await c.post(f"/api/admin/quickbi/assets/{board['id']}/adopt-questions", json={}))
        items = created["items"] if isinstance(created, dict) else created
        assert items and all(q["scenic_id"] == "wys" for q in items)


# ── 对齐 1.2.4:运行记录搜索 / 终止,固定问题筛选 / 思考 / 队列 / 批量运维 ──────────

@run_async
async def test_runs_search_by_question_or_fixed_title_and_abort():
    from cogni_agent.tool_center_service.quickbi.models import QbiRun

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "昨日景区整体运营", "question": "{{昨日}} 各渠道销售额",
            "capability": "dataset_qa", "target": {"cube_id": mock.CUBE_SALES}}))
        run = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run", json={"stream": False}))
        # 运行问句已替换过变量,与标题不是同一段字 —— 按标题搜也要命中,并标出来历
        hit = data(await c.get("/api/admin/quickbi/runs", params={"q": "整体运营"}))["items"]
        assert [r["id"] for r in hit] == [run["id"]]
        assert hit[0]["fixed_question_title"] == "昨日景区整体运营"
        assert data(await c.get("/api/admin/quickbi/runs", params={"q": "各渠道销售"}))["items"]
        assert data(await c.get("/api/admin/quickbi/runs", params={"q": "不存在的问题"}))["items"] == []

        # 僵尸 running(发版/重启掐断留下的):单条终止 + 一键终止
        async with common_db.get_sessionmaker()() as db:
            for rid in ("zombie-1", "zombie-2"):
                db.add(QbiRun(id=rid, connection_id=conn["id"], capability="dataset_qa", mode="direct",
                              question="卡住的", status="running", user_id=ADMIN.user_id))
            await db.commit()
        one = data(await c.post("/api/admin/quickbi/runs/zombie-1/abort"))
        assert one["terminated"] == 1
        again = data(await c.post("/api/admin/quickbi/runs/zombie-1/abort"))
        assert again["terminated"] == 0 and "无需终止" in again["hint"]
        done = data(await c.post("/api/admin/quickbi/runs/abort-batch", json={"stream": False}))
        assert done["total"] == 1 and done["succeeded"] == 1
        z = data(await c.get("/api/admin/quickbi/runs/zombie-2"))
        assert z["status"] == "failed" and z["error_code"] == "aborted"


@run_async
async def test_fixed_list_filters_thinking_queue_and_can_write():
    from cogni_agent.tool_center_service.quickbi.models import QbiRun

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        a = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "解读销售", "question": "解读本月销售",
            "capability": "insight", "target": {"cube_id": mock.CUBE_SALES},
            "prewarm_enabled": True, "prewarm_interval_minutes": 30, "prewarm_ttl_minutes": 60}))
        b = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "客流", "question": "本月客流",
            "capability": "dataset_qa", "target": {"cube_id": mock.CUBE_TRAFFIC}}))

        async def items(params):
            return [x["id"] for x in data(await c.get("/api/admin/quickbi/fixed-questions", params=params))["items"]]
        assert await items({"prewarm": "true"}) == [a["id"]]
        assert await items({"last_status": "none"}) == [a["id"], b["id"]]
        assert await items({"keyword": b["id"]}) == [b["id"]]       # 按问题 ID 搜
        listed = data(await c.get("/api/admin/quickbi/fixed-questions"))["items"]
        assert all(x["can_write"] and x["can_delete"] for x in listed)

        # 查看思考:只认成功且有思考的解读/报告类运行
        assert data(await c.get(f"/api/admin/quickbi/fixed-questions/{a['id']}/thinking"))["found"] is False
        async with common_db.get_sessionmaker()() as db:
            db.add(QbiRun(id="t-qa", connection_id=conn["id"], capability="dataset_qa", mode="fixed",
                          question="q", status="success", user_id=ADMIN.user_id, fixed_question_id=a["id"],
                          result={"thinking": {"reasoning": "选中数据集:销售"}}))
            db.add(QbiRun(id="t-ins", connection_id=conn["id"], capability="insight", mode="fixed",
                          question="解读本月销售", status="success", user_id=ADMIN.user_id,
                          fixed_question_id=a["id"], result={"thinking": {"reasoning": "先看环比"}}))
            await db.commit()
        th = data(await c.get(f"/api/admin/quickbi/fixed-questions/{a['id']}/thinking"))
        assert th["found"] and th["run_id"] == "t-ins" and th["text"] == "先看环比"

        q = data(await c.get("/api/admin/quickbi/fixed-questions/queue-stats"))
        assert {"queued", "running", "prewarm", "others", "oldest_queued_seconds"} <= set(q)


@run_async
async def test_fixed_batch_ops_log_each_row():
    scenic_user = CurrentUser(user_id="u-wys", scenic_id="wys", role="scenic_operator")
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        other = data(await c.post("/api/admin/quickbi/connections", json={
            "name": "九龙潭", "server_domain": "https://bi.aliyun.com", "user_token": "beefbeefbeefbeefbeefbeefbeefbeef",
            "use_trial": True, "scenic_id": "jlt"}))
        data(await c.post("/api/admin/quickbi/assets/datasets", json={
            "connection_id": conn["id"], "cube_id": mock.CUBE_TRAFFIC, "name": "九龙潭-客流-入园明细"}))
        on = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "开了加速", "question": "本月客流", "capability": "dataset_qa",
            "domain": "客流", "prewarm_enabled": True, "prewarm_interval_minutes": 30, "prewarm_ttl_minutes": 60}))
        off = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "没开加速", "question": "本月销售", "capability": "dataset_qa"}))

        r = data(await c.post("/api/admin/quickbi/fixed-questions/batch-op", json={
            "op": "set_interval", "ids": [on["id"], off["id"], "issue_nope"], "interval_minutes": 120,
            "stream": False}))
        assert (r["total"], r["succeeded"], r["failed"]) == (3, 1, 2)
        assert "跳过" in r["items"][1]["message"]
        got = data(await c.get("/api/admin/quickbi/fixed-questions", params={"keyword": on["id"]}))["items"][0]
        # 只改刷新时间时,原存储时间(60)比它短 → 一起抬到 120,否则保存会被「存储<刷新」挡掉
        assert (got["prewarm_interval_minutes"], got["prewarm_ttl_minutes"]) == (120, 120)

        # SSE:逐条回事件
        resp = await c.post("/api/admin/quickbi/fixed-questions/batch-op", json={
            "op": "copy", "ids": [off["id"]], "connection_id": other["id"]})
        evs = sse_events(resp)
        assert [e["type"] for e in evs] == ["start", "item", "done"]
        copied = evs[1]["new_id"]
        cp = data(await c.get("/api/admin/quickbi/fixed-questions", params={"keyword": copied}))["items"][0]
        assert cp["connection_id"] == other["id"] and cp["scenic_id"] == "jlt" and cp["question"] == "本月销售"

        # 解读:按主题(没挂标签时退到数据域)匹配数据集名称
        r = data(await c.post("/api/admin/quickbi/fixed-questions/batch-op", json={
            "op": "insight", "ids": [on["id"], off["id"]], "stream": False}))
        assert [i["ok"] for i in r["items"]] == [True, False]
        got = data(await c.get("/api/admin/quickbi/fixed-questions", params={"keyword": on["id"]}))["items"][0]
        assert got["capability"] == "insight" and got["target"]["cube_id"] == mock.CUBE_TRAFFIC

        # 删除:景区操作员删不了集团级问题,按条失败,其余照删
        async with _client(_app(scenic_user)) as sc:
            r = data(await sc.post("/api/admin/quickbi/fixed-questions/batch-op", json={
                "op": "delete", "ids": [off["id"]], "stream": False}))
            assert r["failed"] == 1
        r = data(await c.post("/api/admin/quickbi/fixed-questions/batch-op", json={
            "op": "delete", "ids": [off["id"], copied], "stream": False}))
        assert r["succeeded"] == 2
        bad = await c.post("/api/admin/quickbi/fixed-questions/batch-op", json={"op": "nope", "ids": [on["id"]]})
        assert bad.status_code == 400


@run_async
async def test_fixed_insight_on_pasted_dashboard_or_portal_url():
    """固定问题的「数据解读」可以直接挂仪表板地址 / 数据门户地址(不必先到数据资产登记)。

    运行时现拉看板结构、按门户菜单解析出 pageId,再拍快照解读;列表上显示看板名称而不是一长串地址。
    """
    portal = "https://bi.aliyun.com/product/view.htm?module=dashboard&productId=p-1"
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        for url in (portal, TOKEN3RD):
            parsed = data(await c.post("/api/admin/quickbi/dashboards/parse",
                                       json={"connection_id": conn["id"], "url": url}))
            assert parsed["page_id"] == mock.PAGE_ID
            fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
                "connection_id": conn["id"], "title": "解读驾驶舱", "question": "解读本月经营情况",
                "capability": "insight",
                "target": {"dashboard": url, "asset_name": parsed.get("name")}}))
            assert fq["target_data_name"] == parsed.get("name")
            assert fq["target_data_id"] == ("p-1" if url == portal else mock.PAGE_ID)
            mock.STATE.__init__()
            run = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run", json={"stream": False}))
            assert run["status"] == "success", run
            paths = [p for _, p in mock.STATE.calls]
            assert "/openapi/v2/snapshot/calling/shot" in paths
            assert ("/openapi/v2/dataportal/query" in paths) == (url == portal)


@run_async
async def test_register_all_datasets_and_clear_skips_referenced():
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        r = data(await c.post("/api/admin/quickbi/assets/datasets/batch",
                              json={"connection_id": conn["id"], "all": True}))
        assert r["created"] >= 2
        again = data(await c.post("/api/admin/quickbi/assets/datasets/batch",
                                  json={"connection_id": conn["id"], "all": True}))
        assert again["created"] == 0 and again["skipped"] == r["created"]     # 已登记的跳过,不堆重复行
        ds = data(await c.get("/api/admin/quickbi/assets", params={"kind": "dataset"}))["items"]
        used = next(d for d in ds if d["ref_id"] == mock.CUBE_SALES)
        data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "销售", "question": "本月销售额", "capability": "dataset_qa",
            "target": {"asset_id": used["id"]}}))
        cleared = data(await c.delete("/api/admin/quickbi/assets/datasets/bulk",
                                      params={"connection_id": conn["id"]}))
        assert cleared["deleted"] == len(ds) - 1 and cleared["skipped"] == 1     # 被问题引用的那个留着
        left = data(await c.get("/api/admin/quickbi/assets", params={"kind": "dataset"}))["items"]
        assert [d["id"] for d in left] == [used["id"]]


# ── 0091:口径说明库 / 景区数据集映射 / 运行记录关联任务号 ───────────────────

@run_async
async def test_prompt_library_crud_and_permissions():
    scenic_user = CurrentUser(user_id="u-wys", scenic_id="wys", role="scenic_admin")
    async with _client(_app()) as c:
        p = data(await c.post("/api/admin/quickbi/prompts", json={
            "name": "同比口径", "content": "同比=(本期-去年同期)/去年同期;缺数据的景区按 0 计", "description": "票务通用"}))
        assert p["visibility"] == "group" and p["can_write"]
        assert [x["id"] for x in data(await c.get("/api/admin/quickbi/prompts", params={"keyword": "同比"}))["items"]] == [p["id"]]
        bad = await c.post("/api/admin/quickbi/prompts", json={"name": "空的", "content": "  "})
        assert bad.status_code == 400
        upd = data(await c.put(f"/api/admin/quickbi/prompts/{p['id']}", json={"content": "", "name": "同比口径v2"}))
        assert upd["name"] == "同比口径v2" and upd["content"].startswith("同比=")     # 空串不清空正文
    async with _client(_app(scenic_user)) as sc:
        items = data(await sc.get("/api/admin/quickbi/prompts"))["items"]
        assert [x["id"] for x in items] == [p["id"]] and items[0]["can_write"] is False   # 集团级共享:可用不可改
        assert (await sc.put(f"/api/admin/quickbi/prompts/{p['id']}", json={"name": "改"})).status_code == 403
        mine = data(await sc.post("/api/admin/quickbi/prompts", json={"name": "我的", "content": "先结论"}))
        assert mine["scenic_id"] == "wys" and mine["visibility"] == "scenic"
    async with _client(_app()) as c:
        data(await c.delete(f"/api/admin/quickbi/prompts/{p['id']}"))
        assert p["id"] not in [x["id"] for x in data(await c.get("/api/admin/quickbi/prompts"))["items"]]


@run_async
async def test_scenic_dataset_map_feeds_default_target_and_connection():
    scenic_user = CurrentUser(user_id="u-wys", scenic_id="wys", role="scenic_admin")
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        m = data(await c.post("/api/admin/quickbi/scenic-datasets", json={
            "scenic_id": "wys", "connection_id": conn["id"], "cube_id": mock.CUBE_TRAFFIC}))
        assert m["cube_name"] and m["prewarm_enabled"] is True and m["connection_name"]   # 名称按实时清单补齐
        # 没指定目标数据、按景区发起 → 用映射的数据集(显式配置,不是猜)
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "客流", "question": "本月客流", "capability": "dataset_qa"}))
        run = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run",
                                json={"scenic_id": "wys", "stream": False}))
        assert run["status"] == "success" and run["target"]["cube_id"] == mock.CUBE_TRAFFIC
        assert run["target"]["from_scenic_map"] is True
        assert mock.STATE.last_body.get("cubeId") == mock.CUBE_TRAFFIC
        # 指定了目标数据就以它为准
        fq2 = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "销售", "question": "本月销售额", "capability": "dataset_qa",
            "target": {"cube_id": mock.CUBE_SALES}}))
        run2 = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq2['id']}/run",
                                 json={"scenic_id": "wys", "stream": False}))
        assert run2["target"]["cube_id"] == mock.CUBE_SALES and not run2["target"].get("from_scenic_map")
        # 问题绑的连接失效时优先换成映射里的那条
        group = svc.Actor(user_id="u-admin", scenic_id=None, role="super_admin", source="console")
        assert await svc.usable_connection_id(group, "gone", "wys") == conn["id"]
        # 景区级预热开关
        off = data(await c.put("/api/admin/quickbi/scenic-datasets/wys/prewarm", json={"enabled": False}))
        assert off["prewarm_enabled"] is False and not await svc.scenic_prewarm_allowed("wys")
        assert await svc.scenic_prewarm_allowed("no-map")
    async with _client(_app(scenic_user)) as sc:
        assert [x["scenic_id"] for x in data(await sc.get("/api/admin/quickbi/scenic-datasets"))["items"]] == ["wys"]
        assert (await sc.post("/api/admin/quickbi/scenic-datasets", json={
            "scenic_id": "wys", "connection_id": conn["id"], "cube_id": "x"})).status_code == 403
    async with _client(_app()) as c:
        data(await c.delete("/api/admin/quickbi/scenic-datasets/wys"))
        assert data(await c.get("/api/admin/quickbi/scenic-datasets"))["items"] == []


@run_async
async def test_runs_carry_the_triggering_task_id():
    from cogni_agent.common.task_context import reset_current_task_id, set_current_task_id

    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        token = set_current_task_id("task-abc")
        try:
            async for _ in svc.run_stream(svc.Actor(user_id=ADMIN.user_id, scenic_id=None, role="super_admin"),
                                          connection_id=conn["id"], capability="dataset_qa",
                                          question="本月客流", target={"cube_id": mock.CUBE_TRAFFIC}):
                pass
        finally:
            reset_current_task_id(token)
        hit = data(await c.get("/api/admin/quickbi/runs", params={"task_id": "task-abc"}))["items"]
        assert len(hit) == 1 and hit[0]["task_id"] == "task-abc"
        assert data(await c.get("/api/admin/quickbi/runs", params={"task_id": "other"}))["items"] == []


@run_async
async def test_register_portal_url_as_dashboard_and_pick_it_in_a_question():
    """固定问题配置里「新增仪表板 / 门户地址」= 当场登记成数据资产再选中:之后别的问题也能从下拉里选到它,
    并且能按图表范围解读。门户来源要回显,下拉里认得出是门户里的哪一页。"""
    portal = "https://bi.aliyun.com/product/view.htm?module=dashboard&productId=p-1"
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        a = data(await c.post("/api/admin/quickbi/assets/dashboards",
                              json={"connection_id": conn["id"], "url": portal, "scenic_id": "wys"}))
        assert a["ref_id"] == mock.PAGE_ID and a["portal"]["portal_name"] == "经营门户"
        assert a["portal"]["menu_title"] == "驾驶舱"
        listed = data(await c.get("/api/admin/quickbi/assets", params={"kind": "dashboard"}))["items"]
        assert any(x["id"] == a["id"] and x["portal"] for x in listed)
        fq = data(await c.post("/api/admin/quickbi/fixed-questions", json={
            "connection_id": conn["id"], "title": "门户解读", "question": "解读本月经营", "capability": "insight",
            "target": {"asset_id": a["id"]}}))
        run = data(await c.post(f"/api/admin/quickbi/fixed-questions/{fq['id']}/run",
                                json={"scenic_id": "wys", "stream": False}))
        assert run["status"] == "success", run


# ─────────────────────────── 文档解析 ───────────────────────────

def test_doc_parse_helpers():
    import openpyxl

    from cogni_agent.tool_center_service.quickbi import docparse as dp

    assert dp.pick_mode("把这些发票提取成台账") == "extract"
    assert dp.pick_mode("这份合同说了什么") == "text"
    assert dp.pick_mode("提取字段", "text") == "text"          # 显式指定优先
    sheets = dp.parse_extraction('好的:\n```json\n{"extraction_data": {"增值税发票": {"headers_cn": '
                                 '["源文件名", "发票号码"], "rows": [["a.pdf", "123"], ["b.pdf", null]]}}}\n```')
    assert sheets == {"增值税发票": {"headers_cn": ["源文件名", "发票号码"], "rows": [["a.pdf", "123"], ["b.pdf", ""]]}}
    with pytest.raises(dp.QuickBIError):
        dp.parse_extraction("抱歉,我无法处理")
    wb = openpyxl.load_workbook(io.BytesIO(dp.build_excel(sheets, scan_time="2026-10-22 10:00:00", total_files=2)))
    assert wb.sheetnames == ["汇总", "增值税发票"]
    ws = wb["增值税发票"]
    assert ws["A1"].value == "源文件名" and ws["A1"].fill.fgColor.rgb.endswith("4472C4") and ws["A1"].font.bold
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert [c.value for c in wb["汇总"][2]] == ["A.财务与税务", "增值税发票", 2, "发票号码"]
    # 本地能读的不走远程;扫描件这类读不出字的返回 None
    assert "1月" in dp.local_text("a.csv", "月份,销售额\n1月,100\n2月,150\n3月,180\n".encode())
    assert dp.local_text("scan.png", b"\x89PNG....") is None


@run_async
async def test_doc_parse_text_mode_remote_ocr_and_local_docx():
    """扫描件 / PPT 走 Quick BI 远程 OCR(upload 只带 file、task 轮询);本地读得出的(表格)不上传;
    一份识别失败不拖累其余文件,失败原因写进结果。"""
    csv_body = "合同,金额\n景区年度采购合同,合同金额人民币壹佰万元整\n".encode()
    mock.STATE.doc_fail_names.add("模糊.png")
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        ids = []
        for name, body in (("发票.pptx", b"PK-fake-pptx"), ("合同.csv", csv_body), ("模糊.png", b"\x89PNG..")):
            ids.append(data(await c.post("/api/admin/quickbi/files", data={"connection_id": conn["id"]},
                                         files={"file": (name, body, "application/octet-stream")}))["id"])
        r = await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "doc_parse", "question": "看看这几份文件写了什么",
            "target": {"file_ids": ids}, "stream": True})
        evs = sse_events(r)
        done = evs[-1]
        assert done["status"] == "success", done
        assert "发票号码 12345678" in done["answer"] and "合同金额人民币壹佰万元整" in done["answer"]
        assert "模糊.png" in done["answer"] and "图片模糊无法识别" in done["answer"]
        uploads = [p for m, p in mock.STATE.calls if p == "/openapi/v2/document/upload"]
        assert len(uploads) == 2                      # csv 本地读,没上传
        assert mock.STATE.doc_polls["doc-task-发票.pptx"] == 2   # 第一次 processing,第二次 success
        run = data(await c.get(f"/api/admin/quickbi/runs/{done['run_id']}"))
        assert run["capability"] == "doc_parse" and run["capability_source"] == "上传文件"


@run_async
async def test_doc_parse_extract_mode_builds_excel_and_chains_to_file_qa(monkeypatch):
    """要求「提取成台账」→ 按分类体系抽字段(大模型)→ 产出一份 xlsx 上传文件,能直接拿它问数。"""
    import openpyxl

    from cogni_agent.runtime_service.llm import client as llm_client

    seen: list = []

    async def fake_llm(model, messages, **kw):
        seen.append(messages)
        return llm_client.LLMResult(content=json.dumps({"extraction_data": {
            "增值税发票": {"headers_cn": ["源文件名", "发票号码", "开票日期", "价税合计"],
                         "rows": [["发票.pptx", "12345678", "2026-08-01", "106.00"]]}}}, ensure_ascii=False))

    monkeypatch.setattr(llm_client, "complete_collected", fake_llm)
    async with _client(_app()) as c:
        conn = await _setup_connection(c)
        fid = data(await c.post("/api/admin/quickbi/files", data={"connection_id": conn["id"]},
                                files={"file": ("发票.pptx", b"PK-fake", "application/octet-stream")}))["id"]
        run = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "doc_parse", "question": "把发票提取成台账",
            "target": {"file_id": fid}, "stream": False}))
        assert run["status"] == "success", run
        assert "增值税发票" in seen[0][0]["content"] and "发票号码 12345678" in seen[0][1]["content"]
        assert run["result"]["charts"][0]["title"] == "增值税发票"
        out = run["result"]["meta"]["output_file"]
        files = data(await c.get("/api/admin/quickbi/files"))["items"]
        xlsx = [f for f in files if f["file_name"].startswith("doc_scan_result_")]
        assert len(xlsx) == 1 and xlsx[0]["id"] in run["answer"]
        assert out["id"] == xlsx[0]["id"]
        from cogni_agent.tool_center_service.quickbi.models import QbiFile
        async with common_db.get_sessionmaker()() as db:
            row = await db.get(QbiFile, xlsx[0]["id"])
        dl = await c.get(f"/api/admin/quickbi/files/{xlsx[0]['id']}/download")
        assert dl.status_code == 200 and dl.content == row.content
        assert "attachment" in dl.headers["content-disposition"]
        wb = openpyxl.load_workbook(io.BytesIO(row.content))
        assert wb.sheetnames == ["汇总", "增值税发票"] and wb["增值税发票"]["B2"].value == "12345678"
        qa = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "file_qa", "question": "一共多少钱",
            "target": {"file_id": xlsx[0]["id"]}, "stream": False}))
        assert qa["status"] == "success", qa


@run_async
async def test_auto_route_documents_to_doc_parse_and_chain(monkeypatch):
    """对齐 SKILL.md:目标里有文档 → 先文档解析;多份文档问具体数据 → 解析成 Excel 后接文件问数;
    显式选了解读、目标却是 PDF → 解析后接解读,而不是报「只支持表格」。"""
    from cogni_agent.runtime_service.llm import client as llm_client

    async def fake_llm(model, messages, **kw):
        return llm_client.LLMResult(content=json.dumps({"extraction_data": {
            "费用报销单": {"headers_cn": ["源文件名", "销售额"], "rows": [["a.png", "106.00"], ["b.png", "50.00"]]}}},
            ensure_ascii=False))

    monkeypatch.setattr(llm_client, "complete_collected", fake_llm)
    async with _client(_app()) as c:
        conn = await _setup_connection(c)

        async def up(name):
            return data(await c.post("/api/admin/quickbi/files", data={"connection_id": conn["id"]},
                                     files={"file": (name, b"\x89PNG..", "image/png")}))["id"]
        a, b = await up("a.png"), await up("b.png")
        single = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "question": "这张图上写了什么", "target": {"file_id": a}, "stream": False}))
        assert single["capability"] == "doc_parse" and "doc_then" not in single["target"]
        assert "发票号码 12345678" in single["answer"]
        multi = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "question": "各部门人数TOP", "target": {"file_ids": [a, b]},
            "stream": False}))
        assert multi["status"] == "success", multi
        assert multi["capability"] == "doc_parse" and multi["target"]["doc_then"] == "file_qa"
        assert "研发部人数最多" in multi["answer"]                   # 文件问数(对解析出的 Excel)的回答
        assert ("POST", "/openapi/v2/smartq/queryByQuestionStreamByFile") in mock.STATE.calls
        ins = data(await c.post("/api/admin/quickbi/run", json={
            "connection_id": conn["id"], "capability": "insight", "question": "解读一下",
            "target": {"file_id": a}, "stream": False}))
        assert ins["status"] == "success", (ins["error"], ins["error_code"])
        assert ins["capability"] == "doc_parse" and ins["target"]["doc_then"] == "insight" and "环比" in ins["answer"]
